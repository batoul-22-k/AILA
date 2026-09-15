from __future__ import annotations

import json
import os
import pandas as pd
import requests
import tempfile
import unittest
import zipfile
from importlib.util import find_spec
from pathlib import Path
from unittest.mock import call, patch

from app.llm_evaluation.calculate_results import calculate_per_model, validate_ranges
from app.llm_evaluation.calibration_benchmark import (
    parse_calibration_pairs,
    rank_rows,
    run_calibration,
)
from app.llm_evaluation.config import EvaluationConfig, load_config
from app.llm_evaluation.diagnose_model import DIAGNOSTIC_PROMPT, run_rq1_diagnostic
from app.llm_evaluation.export_human_review import build_review_rows
from app.llm_evaluation.merge_model_runs import merge_payloads
from app.llm_evaluation.prompts import OLLAMA_JSON_SCHEMA, ONE_QUESTION_JSON_SCHEMA, build_one_question_prompt
from app.llm_evaluation.run_comparison import (
    aggregate_one_question_records,
    build_generate_payload,
    call_ollama,
    load_passages,
    make_one_question_record,
    make_record,
    passage_hash,
    parse_json_output,
    perform_generate_request,
    prompt_size_diagnostics,
    run_experiment,
    summarize_records,
    wait_until_models_absent,
)
from app.llm_evaluation.validate_outputs import (
    analyze_model_records,
    calculate_metrics,
    complete_question_set,
    final_output_records,
    mcq_has_four_options,
    one_question_schema_diagnostics,
    questions_from_parsed,
    relevance_flags,
    schema_compliant,
    schema_diagnostics,
    unsupported_question_terms,
    valid_answer_index,
    valid_bloom,
    valid_difficulty,
)


def valid_mcq(index: int = 1) -> dict:
    return {
        "q": f"Which concept reduces latency {index}?",
        "o": ["Caching resources", "Large transfers", "Distant servers", "Delayed feedback"],
        "a": 0,
        "e": "Caching is supported by the lecture.",
        "b": "Understand",
        "d": "medium",
        "s": 1,
    }


def valid_one_question_output(index: int = 1) -> dict:
    return {"question": valid_mcq(index)}


def child_question_record(index: int = 1, *, success: bool = True, truncated: bool = False, elapsed: float = 0.1) -> dict:
    parsed = valid_one_question_output(index) if success and not truncated else {"question": valid_mcq(index)}
    return {
        "experiment_id": "exp",
        "experiment_group_id": "group",
        "record_type": "individual_question",
        "request_id": f"req_{index}",
        "lecture_id": "lecture_1",
        "attempt_number": 1,
        "question_sequence_number": index,
        "model": "tinyllama",
        "request_success": success,
        "generation_completed": success and not truncated,
        "generation_success": success and not truncated,
        "valid_json": success,
        "schema_compliant": success,
        "done": success,
        "done_reason": "length" if truncated else "stop",
        "truncated": truncated,
        "invalid_control_characters": None,
        "elapsed_seconds": elapsed,
        "http_status": 200 if success else 500,
        "parsed_output": parsed,
        "prompt_hash": f"prompt_{index}",
        "error": "failed" if (not success or truncated) else None,
        "error_category": "truncated_output" if truncated else (None if success else "http_failure"),
    }


def test_config(
    use_json_schema: bool = True,
    request_delay_seconds: int = 0,
    num_ctx: int = 2048,
    num_predict: int = 512,
    questions_per_request: int = 4,
) -> EvaluationConfig:
    return EvaluationConfig(
        model_a="tinyllama",
        model_b="llama3.1:8b",
        ollama_base_url="http://localhost:11434",
        timeout_seconds=1,
        prompt_chars=1800,
        num_ctx=num_ctx,
        num_predict=num_predict,
        temperature=0,
        seed=42,
        keep_alive=0,
        unload_timeout_seconds=1,
        request_delay_seconds=request_delay_seconds,
        use_json_schema=use_json_schema,
        questions_per_request=questions_per_request,
    )


class FakeResponse:
    def __init__(self, status_code: int = 200, payload: dict | None = None, text: str = "") -> None:
        self.status_code = status_code
        self._payload = payload
        self.text = text

    def json(self) -> dict:
        if self._payload is None:
            raise json.JSONDecodeError("bad json", self.text, 0)
        return self._payload

    def raise_for_status(self) -> None:
        if self.status_code >= 400:
            error = requests.HTTPError(f"HTTP {self.status_code}")
            error.response = self
            raise error


class EvaluationModuleTests(unittest.TestCase):
    def test_malformed_json_is_not_valid(self) -> None:
        valid, parsed = parse_json_output('{"questions": [')
        self.assertFalse(valid)
        self.assertIsNone(parsed)

    def test_default_generation_limits_are_increased(self) -> None:
        with patch.dict(os.environ, {}, clear=True):
            config = load_config()
        self.assertEqual(config.num_ctx, 2048)
        self.assertEqual(config.num_predict, 700)
        self.assertEqual(config.temperature, 1)
        self.assertEqual(config.keep_alive, "2m")
        self.assertFalse(config.use_json_schema)
        self.assertEqual(config.models, ("tinyllama-local", "gpt-oss-20b-cloud"))
        self.assertEqual(config.model_tag("tinyllama-local"), "tinyllama")
        self.assertEqual(config.model_config("gpt-oss-20b-cloud")["deployment_type"], "cloud")
        self.assertEqual(config.model_config("gpt-oss-20b-cloud")["pricing"]["direct_request_cost_usd"], None)

    def test_generation_limit_environment_overrides(self) -> None:
        with patch.dict(os.environ, {"EVALUATION_NUM_CTX": "3072", "EVALUATION_NUM_PREDICT": "640"}, clear=True):
            config = load_config()
        self.assertEqual(config.num_ctx, 3072)
        self.assertEqual(config.num_predict, 640)

    def test_questions_per_request_environment_override(self) -> None:
        with patch.dict(os.environ, {"EVALUATION_QUESTIONS_PER_REQUEST": "1"}, clear=True):
            config = load_config()
        self.assertEqual(config.questions_per_request, 1)

    def test_generation_limit_environment_overrides_must_be_positive(self) -> None:
        with patch.dict(os.environ, {"EVALUATION_NUM_CTX": "0"}, clear=True):
            with self.assertRaises(ValueError):
                load_config()

    def test_missing_questions_returns_empty_list(self) -> None:
        self.assertEqual(questions_from_parsed({"not_questions": []}), [])

    def test_one_question_prompt_and_schema(self) -> None:
        prompt = build_one_question_prompt("Slide 1: Routing reduces path delay.", 1200, 2)
        self.assertIn("exactly one key named \"question\"", prompt)
        self.assertIn("question sequence 2 of 4", prompt)
        diagnostics = one_question_schema_diagnostics(valid_one_question_output(1))
        self.assertTrue(diagnostics["compliant"])
        bad = one_question_schema_diagnostics({"questions": [valid_mcq(1)]})
        self.assertEqual(bad["reason"], "missing_question_key")

    def test_one_question_schema_payload_can_be_requested(self) -> None:
        payload = build_generate_payload(test_config(use_json_schema=True), "tinyllama", "prompt", output_schema=ONE_QUESTION_JSON_SCHEMA)
        self.assertEqual(payload["format"], ONE_QUESTION_JSON_SCHEMA)

    @unittest.skipIf(find_spec("pptx") is None, "python-pptx is not installed")
    def test_load_passages_accepts_pptx_dataset_directory(self) -> None:
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as temp_dir:
            dataset_dir = Path(temp_dir)
            deck_path = dataset_dir / "Course One.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = "Routing overview"
            presentation.save(deck_path)

            passages = load_passages(dataset_dir)

        self.assertEqual(len(passages), 1)
        self.assertEqual(passages[0]["id"], "course_one")
        self.assertEqual(passages[0]["topic"], "Course One")
        self.assertEqual(passages[0]["source_type"], "pptx")
        self.assertIn("[Slide 1]", passages[0]["content"])
        self.assertIn("Routing overview", passages[0]["content"])

    def test_load_passages_accepts_pptx_directory_without_python_pptx(self) -> None:
        slide_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
       xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>TCP congestion control</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld>
</p:sld>"""
        with tempfile.TemporaryDirectory() as temp_dir:
            dataset_dir = Path(temp_dir)
            deck_path = dataset_dir / "TCP.pptx"
            with zipfile.ZipFile(deck_path, "w") as archive:
                archive.writestr("ppt/slides/slide1.xml", slide_xml)
            with patch.dict("sys.modules", {"pptx": None}):
                passages = load_passages(dataset_dir)

        self.assertEqual(len(passages), 1)
        self.assertEqual(passages[0]["id"], "tcp")
        self.assertIn("[Slide 1]", passages[0]["content"])
        self.assertIn("TCP congestion control", passages[0]["content"])

    def test_incorrect_question_count_is_incomplete(self) -> None:
        questions = [valid_mcq(), valid_mcq(2), valid_mcq(3)]
        self.assertFalse(complete_question_set(questions))
        self.assertEqual(schema_diagnostics({"questions": questions})["reason"], "invalid_question_count")

    def test_valid_canonical_object_root_is_complete(self) -> None:
        output = {"questions": [valid_mcq(1), valid_mcq(2), valid_mcq(3), valid_mcq(4)]}
        self.assertTrue(schema_compliant(output))
        self.assertTrue(complete_question_set(questions_from_parsed(output)))

    def test_top_level_array_is_rejected(self) -> None:
        output = [valid_mcq(1), valid_mcq(2), valid_mcq(3), valid_mcq(4)]
        diagnostics = schema_diagnostics(output)
        self.assertFalse(schema_compliant(output))
        self.assertEqual(questions_from_parsed(output), [])
        self.assertEqual(diagnostics["reason"], "root_not_object")

    def test_missing_questions_key_is_rejected(self) -> None:
        diagnostics = schema_diagnostics({"items": [valid_mcq(1)]})
        self.assertFalse(diagnostics["compliant"])
        self.assertEqual(diagnostics["reason"], "missing_questions_key")

    def test_extra_root_keys_are_rejected(self) -> None:
        diagnostics = schema_diagnostics({"questions": [valid_mcq(1), valid_mcq(2), valid_mcq(3), valid_mcq(4)], "meta": {}})
        self.assertFalse(diagnostics["compliant"])
        self.assertEqual(diagnostics["reason"], "unexpected_root_fields")

    def test_questions_must_be_an_array(self) -> None:
        diagnostics = schema_diagnostics({"questions": {"q": "Which concept reduces latency?"}})
        self.assertFalse(diagnostics["compliant"])
        self.assertEqual(diagnostics["reason"], "questions_not_array")

    def test_compact_output_with_excessive_text_length_fails(self) -> None:
        output = {"questions": [valid_mcq(1), valid_mcq(2), valid_mcq(3), valid_mcq(4)]}
        output["questions"][0]["q"] = "Which classroom engagement question asks students to compare network latency bandwidth caching media transfers today?"
        diagnostics = schema_diagnostics(output)
        self.assertFalse(schema_compliant(output))
        self.assertEqual(diagnostics["reason"], "text_length_exceeded")

    def test_invalid_answer_index_fails(self) -> None:
        question = valid_mcq()
        question["a"] = 4
        self.assertFalse(valid_answer_index(question))
        self.assertFalse(schema_compliant({"questions": [question]}))
        self.assertEqual(schema_diagnostics({"questions": [question, valid_mcq(2), valid_mcq(3), valid_mcq(4)]})["reason"], "invalid_answer_index")

    def test_mcq_with_fewer_than_four_options_fails(self) -> None:
        question = valid_mcq()
        question["o"] = ["A", "B", "C"]
        self.assertFalse(mcq_has_four_options(question))
        self.assertFalse(schema_compliant({"questions": [question]}))
        self.assertEqual(schema_diagnostics({"questions": [question, valid_mcq(2), valid_mcq(3), valid_mcq(4)]})["reason"], "invalid_question_fields")

    def test_invalid_bloom_label_fails(self) -> None:
        question = valid_mcq()
        question["b"] = "Comprehend"
        self.assertFalse(valid_bloom(question))

    def test_invalid_difficulty_label_fails(self) -> None:
        question = valid_mcq()
        question["d"] = "Medium"
        self.assertFalse(valid_difficulty(question))

    def test_missing_compact_field_fails(self) -> None:
        question = valid_mcq()
        del question["s"]
        self.assertFalse(schema_compliant({"questions": [question]}))
        self.assertFalse(complete_question_set([question, valid_mcq(2), valid_mcq(3), valid_mcq(4)]))
        self.assertEqual(schema_diagnostics({"questions": [question, valid_mcq(2), valid_mcq(3), valid_mcq(4)]})["reason"], "invalid_question_fields")

    def test_off_topic_question_detection_flags_unsupported_terms(self) -> None:
        lecture = "Latency delays quiz submissions. Caching resources reduces latency."
        question = valid_mcq()
        question["q"] = "Which photosynthesis step reduces chlorophyll?"
        self.assertEqual(unsupported_question_terms(question, lecture), ["chlorophyll", "photosynthesis", "step"])
        flags = relevance_flags([question], lecture)
        self.assertEqual(flags[0]["question_number"], 1)
        self.assertIn("photosynthesis", flags[0]["unsupported_terms"])

    def test_failed_ollama_request_is_recorded(self) -> None:
        def raise_connection_error(*args, **kwargs):
            raise requests.ConnectionError("offline")

        config = test_config()
        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.wait_until_models_absent", return_value={"confirmed": True, "running_models": [], "error": None}),
            patch("app.llm_evaluation.run_comparison.requests.post", raise_connection_error),
        ):
            result = call_ollama(config=config, model="tinyllama", prompt="prompt")
        self.assertIsNone(result["raw_response"])
        self.assertIsNotNone(result["error"])
        self.assertIn("failed", str(result["error"]))
        self.assertGreaterEqual(result["elapsed_seconds"], 0)

    def test_keep_alive_zero_is_in_generation_payload(self) -> None:
        payload = build_generate_payload(test_config(), "tinyllama", "prompt")
        self.assertEqual(payload["keep_alive"], 0)
        self.assertNotIn("keep_alive", payload["options"])

    def test_loaded_config_payload_matches_production_generation_defaults(self) -> None:
        with patch.dict(os.environ, {}, clear=True):
            config = load_config()
        payload = build_generate_payload(config, "tinyllama", "prompt")
        self.assertEqual(payload["format"], "json")
        self.assertFalse(payload["stream"])
        self.assertEqual(payload["keep_alive"], "2m")
        self.assertEqual(payload["options"]["temperature"], 1)
        self.assertEqual(payload["options"]["num_ctx"], 2048)
        self.assertEqual(payload["options"]["num_predict"], 700)
        self.assertEqual(payload["options"]["num_gpu"], 0)
        self.assertNotIn("seed", payload["options"])

    def test_generation_settings_record_keep_alive(self) -> None:
        self.assertEqual(test_config().generation_settings["keep_alive"], 0)

    def test_unload_runs_in_finally_after_failed_request(self) -> None:
        config = EvaluationConfig(
            model_a="tinyllama",
            model_b="llama3.1:8b",
            ollama_base_url="http://localhost:11434",
            timeout_seconds=1,
            prompt_chars=1800,
            num_ctx=2048,
            num_predict=512,
            temperature=0,
            seed=42,
            keep_alive=0,
            unload_timeout_seconds=1,
            request_delay_seconds=0,
            use_json_schema=True,
            unload_after_request=True,
        )
        with (
            patch("app.llm_evaluation.run_comparison.requests.post", side_effect=requests.ConnectionError("offline")),
            patch("app.llm_evaluation.run_comparison.unload_model") as unload_model,
        ):
            result = perform_generate_request(config, "tinyllama", build_generate_payload(config, "tinyllama", "prompt"))
        unload_model.assert_called_once_with(config, "tinyllama")
        self.assertEqual(result["exception_type"], "ConnectionError")

    def test_call_ollama_does_not_wait_for_absent_model_when_keep_alive_enabled(self) -> None:
        config = EvaluationConfig(
            model_a="tinyllama",
            model_b="",
            model_ids=("tinyllama",),
            ollama_base_url="http://localhost:11434",
            timeout_seconds=1,
            request_delay_seconds=0,
            unload_after_request=False,
        )
        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", side_effect=[([], None), (["tinyllama:latest"], None)]),
            patch("app.llm_evaluation.run_comparison.wait_until_models_absent") as wait_mock,
            patch("app.llm_evaluation.run_comparison.perform_generate_request", return_value={"raw_response": {"response": "{}", "done": True}, "error": None, "elapsed_seconds": 0.1, "http_status": 200}),
        ):
            result = call_ollama(config=config, model="tinyllama", prompt="prompt")
        wait_mock.assert_not_called()
        self.assertFalse(result["model_unloaded_after_request"])
        self.assertEqual(result["running_models_after_request"], ["tinyllama:latest"])

    def test_wait_until_models_absent_polls_api_ps(self) -> None:
        config = test_config()
        responses = [
            FakeResponse(payload={"models": [{"name": "tinyllama:latest"}]}),
            FakeResponse(payload={"models": []}),
        ]
        with (
            patch("app.llm_evaluation.run_comparison.requests.get", side_effect=responses) as get_mock,
            patch("app.llm_evaluation.run_comparison.time.sleep"),
        ):
            result = wait_until_models_absent(config, ["tinyllama"])
        self.assertTrue(result["confirmed"])
        self.assertEqual(get_mock.call_count, 2)
        self.assertEqual(get_mock.call_args_list[0], call(config.ps_url, timeout=config.timeout_seconds))

    def test_call_ollama_skips_when_other_model_cannot_unload(self) -> None:
        config = test_config()
        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=(["llama3.1:8b"], None)),
            patch("app.llm_evaluation.run_comparison.unload_and_confirm", return_value={"confirmed": False, "running_models": ["llama3.1:8b"], "error": "still loaded"}),
            patch("app.llm_evaluation.run_comparison.perform_generate_request") as generate,
        ):
            result = call_ollama(config=config, model="tinyllama", prompt="prompt", other_models=["llama3.1:8b"])
        generate.assert_not_called()
        self.assertEqual(result["attempt_count"], 0)
        self.assertIn("Skipped", result["error"])

    def test_run_experiment_uses_model_batched_execution(self) -> None:
        config = test_config()
        passages = [
            {"id": "lecture_1", "topic": "A", "content": "Content A"},
            {"id": "lecture_2", "topic": "B", "content": "Content B"},
        ]

        def fake_record(*, experiment_id, experiment_group_id, lecture, model, config, prompt, lecture_hash, other_models, attempt_number, **kwargs):
            return {
                "experiment_id": experiment_id,
                "experiment_group_id": experiment_group_id,
                "lecture_id": lecture["id"],
                "lecture_hash": lecture_hash,
                "topic": lecture["topic"],
                "lecture_content": lecture["content"],
                "model": model,
                "attempt_number": attempt_number,
                "generation_settings": config.generation_settings,
                "elapsed_seconds": 0,
                "valid_json": True,
                "schema_compliant": True,
                "generation_success": True,
                "raw_output": "{}",
                "parsed_output": {},
                "error": None,
            }

        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "inputs.json"
            results_dir = Path(temp_dir) / "results"
            input_path.write_text(json.dumps(passages), encoding="utf-8")
            with (
                patch("app.llm_evaluation.run_comparison.preflight_checks", return_value={}),
                patch("app.llm_evaluation.run_comparison.unload_and_confirm", return_value={"confirmed": True, "running_models": [], "error": None}),
                patch("app.llm_evaluation.run_comparison.make_record", side_effect=fake_record),
            ):
                output_path = run_experiment(input_path, results_dir, config)
            payload = json.loads(output_path.read_text(encoding="utf-8"))
        self.assertEqual(
            [record["model"] for record in payload["records"]],
            ["tinyllama"] * 6 + ["llama3.1:8b"] * 6,
        )
        self.assertEqual([record["attempt_number"] for record in payload["records"][:3]], [1, 2, 3])
        self.assertEqual([entry["model"] for entry in payload["execution_order"]], ["tinyllama", "llama3.1:8b"])

    def test_single_model_execution_runs_only_selected_model(self) -> None:
        config = test_config()
        passages = [
            {"id": "lecture_1", "topic": "A", "content": "Content A"},
            {"id": "lecture_2", "topic": "B", "content": "Content B"},
        ]
        seen_models: list[str] = []

        def fake_record(*, experiment_id, experiment_group_id, lecture, model, config, prompt, lecture_hash, other_models, attempt_number, **kwargs):
            seen_models.append(model)
            return {
                "experiment_id": experiment_id,
                "experiment_group_id": experiment_group_id,
                "lecture_id": lecture["id"],
                "lecture_hash": lecture_hash,
                "model": model,
                "attempt_number": attempt_number,
                "request_success": True,
                "generation_success": True,
                "valid_json": True,
                "schema_compliant": True,
                "complete_question_set": True,
                "empty_response": False,
                "truncated": False,
                "retry_performed": False,
                "generation_settings": config.generation_settings,
                "error": None,
            }

        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "inputs.json"
            results_dir = Path(temp_dir) / "results"
            input_path.write_text(json.dumps(passages), encoding="utf-8")
            with (
                patch("app.llm_evaluation.run_comparison.preflight_checks", return_value={}),
                patch("app.llm_evaluation.run_comparison.unload_and_confirm", return_value={"confirmed": True, "running_models": [], "error": None}),
                patch("app.llm_evaluation.run_comparison.make_record", side_effect=fake_record),
            ):
                output_path = run_experiment(
                    input_path,
                    results_dir,
                    config,
                    selected_model="tinyllama",
                    experiment_group_id="group_1",
                )
            payload = json.loads(output_path.read_text(encoding="utf-8"))

        self.assertEqual(seen_models, ["tinyllama"] * 6)
        self.assertEqual(payload["experiment_group_id"], "group_1")
        self.assertEqual(payload["models"], ["tinyllama"])
        self.assertTrue(payload["single_model_run"])
        self.assertIn("tinyllama", output_path.name)

    def test_four_successful_one_question_requests_create_valid_aggregate(self) -> None:
        config = test_config(questions_per_request=1)
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        records = [
            child_question_record(3, elapsed=0.3),
            child_question_record(1, elapsed=0.1),
            child_question_record(4, elapsed=0.4),
            child_question_record(2, elapsed=0.2),
        ]
        aggregate = aggregate_one_question_records(
            experiment_id="exp",
            experiment_group_id="group",
            lecture=lecture,
            model="tinyllama",
            config=config,
            attempt_number=1,
            question_records=records,
            lecture_hash=passage_hash(lecture),
            input_hash_value="input_hash",
        )
        self.assertTrue(aggregate["generation_success"])
        self.assertTrue(aggregate["complete_question_set"])
        self.assertEqual(aggregate["record_type"], "aggregate_question_set")
        self.assertEqual(aggregate["aggregate_request_ids"], ["req_1", "req_2", "req_3", "req_4"])
        self.assertEqual([question["q"] for question in aggregate["parsed_output"]["questions"]], [valid_mcq(1)["q"], valid_mcq(2)["q"], valid_mcq(3)["q"], valid_mcq(4)["q"]])
        self.assertEqual(aggregate["total_elapsed_seconds"], 1.0)

    def test_one_failed_request_makes_aggregate_fail(self) -> None:
        config = test_config(questions_per_request=1)
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        records = [child_question_record(1), child_question_record(2, success=False), child_question_record(3), child_question_record(4)]
        aggregate = aggregate_one_question_records(
            experiment_id="exp",
            experiment_group_id="group",
            lecture=lecture,
            model="tinyllama",
            config=config,
            attempt_number=1,
            question_records=records,
            lecture_hash=passage_hash(lecture),
            input_hash_value="input_hash",
        )
        self.assertFalse(aggregate["generation_success"])
        self.assertFalse(aggregate["complete_question_set"])
        self.assertEqual(aggregate["successful_question_count"], 3)

    def test_one_truncated_request_makes_aggregate_fail(self) -> None:
        config = test_config(questions_per_request=1)
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        records = [child_question_record(1), child_question_record(2), child_question_record(3, truncated=True), child_question_record(4)]
        aggregate = aggregate_one_question_records(
            experiment_id="exp",
            experiment_group_id="group",
            lecture=lecture,
            model="tinyllama",
            config=config,
            attempt_number=1,
            question_records=records,
            lecture_hash=passage_hash(lecture),
            input_hash_value="input_hash",
        )
        self.assertFalse(aggregate["generation_success"])
        self.assertTrue(aggregate["truncated"])
        self.assertEqual(aggregate["error_category"], "truncated_output")

    def test_one_question_mode_uses_identical_request_counts_across_models(self) -> None:
        config = test_config(questions_per_request=1)
        passages = [{"id": "lecture_1", "topic": "A", "content": "Content A"}]
        calls_by_model: dict[str, int] = {}

        def fake_one_question_record(*, experiment_id, experiment_group_id, lecture, model, config, question_sequence_number, attempt_number, **kwargs):
            calls_by_model[model] = calls_by_model.get(model, 0) + 1
            return {**child_question_record(question_sequence_number), "model": model, "experiment_id": experiment_id, "experiment_group_id": experiment_group_id}

        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "inputs.json"
            results_dir = Path(temp_dir) / "results"
            input_path.write_text(json.dumps(passages), encoding="utf-8")
            with (
                patch("app.llm_evaluation.run_comparison.preflight_checks", return_value={}),
                patch("app.llm_evaluation.run_comparison.unload_and_confirm", return_value={"confirmed": True, "running_models": [], "error": None}),
                patch("app.llm_evaluation.run_comparison.make_one_question_record", side_effect=fake_one_question_record),
            ):
                output_path = run_experiment(input_path, results_dir, config)
            payload = json.loads(output_path.read_text(encoding="utf-8"))
        self.assertEqual(calls_by_model, {"tinyllama": 12, "llama3.1:8b": 12})
        aggregates = [record for record in payload["records"] if record.get("record_type") == "aggregate_question_set"]
        self.assertEqual(len(aggregates), 6)
        self.assertEqual([entry["requests_per_lecture_attempt"] for entry in payload["execution_order"]], [4, 4])

    def make_merge_payload(self, model: str, settings: dict | None = None, hashes: dict | None = None) -> dict:
        settings = settings or test_config().generation_settings
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        hashes = hashes or {"lecture_1": passage_hash(lecture)}
        return {
            "experiment_id": f"exp_{model}",
            "experiment_group_id": "group_1",
            "models": [model],
            "configured_models": ["tinyllama", "llama3.1:8b"],
            "generation_settings": settings,
            "input_passage_hashes": hashes,
            "records": [
                {
                    "experiment_id": f"exp_{model}",
                    "experiment_group_id": "group_1",
                    "lecture_id": "lecture_1",
                    "lecture_hash": hashes["lecture_1"],
                    "model": model,
                    "request_success": True,
                    "generation_success": True,
                    "valid_json": True,
                    "schema_compliant": True,
                    "complete_question_set": True,
                    "empty_response": False,
                    "truncated": False,
                    "retry_performed": False,
                }
            ],
        }

    def test_merge_rejects_generation_settings_mismatch(self) -> None:
        left = self.make_merge_payload("tinyllama")
        changed_settings = dict(test_config().generation_settings)
        changed_settings["num_predict"] = 640
        right = self.make_merge_payload("llama3.1:8b", settings=changed_settings)
        with self.assertRaisesRegex(ValueError, "generation settings differ"):
            merge_payloads([left, right])

    def test_merge_rejects_input_passage_hash_mismatch(self) -> None:
        left = self.make_merge_payload("tinyllama")
        right = self.make_merge_payload("llama3.1:8b", hashes={"lecture_1": "different"})
        with self.assertRaisesRegex(ValueError, "input passage hashes differ"):
            merge_payloads([left, right])

    def test_merge_combines_single_model_runs_successfully(self) -> None:
        left = self.make_merge_payload("tinyllama")
        right = self.make_merge_payload("llama3.1:8b")
        merged = merge_payloads([left, right])
        self.assertEqual(merged["experiment_group_id"], "group_1")
        self.assertEqual(merged["models"], ["tinyllama", "llama3.1:8b"])
        self.assertEqual([record["model"] for record in merged["records"]], ["tinyllama", "llama3.1:8b"])
        self.assertEqual(merged["summary"]["overall"]["total_requests"], 2)
        self.assertEqual(merged["summary"]["overall"]["generation_successes"], 2)

    def test_http_500_retries_identical_request_once(self) -> None:
        config = test_config()
        payloads: list[dict] = []

        def capture_post(url, json, timeout):
            payloads.append(dict(json))
            if len(payloads) == 1:
                return FakeResponse(status_code=500, payload={"error": "runner terminated"}, text="runner terminated")
            return FakeResponse(payload={"response": '{"questions": []}'}, text='{"response": "{\\"questions\\": []}"}')

        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", side_effect=capture_post),
            patch("app.llm_evaluation.run_comparison.time.sleep"),
        ):
            result = call_ollama(config=config, model="tinyllama", prompt="prompt")
        self.assertEqual(result["attempt_count"], 2)
        self.assertTrue(result["retry_performed"])
        self.assertIn("HTTP 500", result["initial_error"])
        self.assertIsNone(result["final_error"])
        self.assertEqual(payloads[0], payloads[1])

    def test_malformed_model_output_does_not_retry(self) -> None:
        config = test_config()
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        payloads: list[dict] = []

        def capture_post(url, json, timeout):
            payloads.append(dict(json))
            return FakeResponse(payload={"response": '{"questions": ['}, text="ok")

        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", side_effect=capture_post),
        ):
            record = make_record(experiment_id="exp", lecture=lecture, model="tinyllama", config=config, prompt="prompt")
        self.assertFalse(record["valid_json"])
        self.assertEqual(record["attempt_count"], 1)
        self.assertFalse(record["retry_performed"])
        self.assertEqual(len(payloads), 1)

    def test_json_schema_mode_enabled_and_disabled(self) -> None:
        enabled_payload = build_generate_payload(test_config(use_json_schema=True), "tinyllama", "prompt")
        disabled_payload = build_generate_payload(test_config(use_json_schema=False), "tinyllama", "prompt")
        self.assertEqual(enabled_payload["format"], OLLAMA_JSON_SCHEMA)
        self.assertEqual(enabled_payload["format"]["type"], "object")
        self.assertEqual(enabled_payload["format"]["required"], ["questions"])
        self.assertFalse(enabled_payload["format"]["additionalProperties"])
        question_schema = enabled_payload["format"]["properties"]["questions"]["items"]
        self.assertFalse(question_schema["additionalProperties"])
        self.assertEqual(disabled_payload["format"], "json")
        self.assertEqual(test_config(use_json_schema=True).generation_settings["format"], "json_schema")
        self.assertTrue(test_config(use_json_schema=True).generation_settings["use_json_schema"])
        self.assertFalse(test_config(use_json_schema=False).generation_settings["use_json_schema"])

    def test_generation_settings_are_identical_for_both_models(self) -> None:
        config = test_config()
        tiny_payload = build_generate_payload(config, "tinyllama", "prompt")
        llama_payload = build_generate_payload(config, "llama3.1:8b", "prompt")
        self.assertEqual(tiny_payload["prompt"], llama_payload["prompt"])
        self.assertEqual(tiny_payload["stream"], llama_payload["stream"])
        self.assertEqual(tiny_payload["keep_alive"], llama_payload["keep_alive"])
        self.assertEqual(tiny_payload["format"], llama_payload["format"])
        self.assertEqual(tiny_payload["options"], llama_payload["options"])
        self.assertEqual(tiny_payload["options"]["num_ctx"], 2048)
        self.assertEqual(tiny_payload["options"]["num_predict"], 512)

    def test_empty_http_200_response_classification(self) -> None:
        config = test_config()
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", return_value=FakeResponse(payload={"response": "", "done": True}, text="ok")),
        ):
            record = make_record(experiment_id="exp", lecture=lecture, model="tinyllama", config=config, prompt="prompt")
        self.assertFalse(record["generation_completed"])
        self.assertTrue(record["empty_response"])
        self.assertEqual(record["error_category"], "empty_response")
        self.assertEqual(record["error"], "Ollama returned an empty model response")
        self.assertFalse(record["valid_json"])
        self.assertIsNone(record["parsed_output"])
        self.assertEqual(record["raw_response"], {"response": "", "done": True})

    def test_done_false_response_is_incomplete_failure(self) -> None:
        config = test_config()
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", return_value=FakeResponse(payload={"response": "partial", "done": False}, text="ok")),
        ):
            record = make_record(experiment_id="exp", lecture=lecture, model="tinyllama", config=config, prompt="prompt")
        self.assertFalse(record["generation_success"])
        self.assertFalse(record["generation_completed"])
        self.assertEqual(record["error_category"], "incomplete_response")
        self.assertEqual(record["retry_reason"], "done_false")
        self.assertEqual(record["format_validation"]["done"], False)

    def test_control_character_response_is_rejected(self) -> None:
        config = test_config()
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", return_value=FakeResponse(payload={"response": "\x1c\x1c", "done": True}, text="ok")),
        ):
            record = make_record(experiment_id="exp", lecture=lecture, model="tinyllama", config=config, prompt="prompt")
        self.assertFalse(record["generation_success"])
        self.assertEqual(record["error_category"], "invalid_control_characters")
        self.assertEqual(record["invalid_control_characters"]["codepoints"], {"U+001C": 2})
        self.assertEqual(record["format_validation"]["invalid_control_characters"]["count"], 2)

    def test_whitespace_only_http_200_response_classification(self) -> None:
        config = test_config()
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", return_value=FakeResponse(payload={"response": "   \n", "done": True}, text="ok")),
        ):
            record = make_record(experiment_id="exp", lecture=lecture, model="tinyllama", config=config, prompt="prompt")
        self.assertTrue(record["empty_response"])
        self.assertEqual(record["error_category"], "empty_response")

    def test_done_reason_length_truncation_classification(self) -> None:
        config = test_config()
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        partial = '{"questions": [{"type": "mcq"'
        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", return_value=FakeResponse(payload={"response": partial, "done": True, "done_reason": "length", "eval_count": 512}, text="ok")),
        ):
            record = make_record(experiment_id="exp", lecture=lecture, model="llama3.1:8b", config=config, prompt="prompt")
        self.assertTrue(record["truncated"])
        self.assertFalse(record["generation_completed"])
        self.assertEqual(record["error_category"], "truncated_output")
        self.assertEqual(record["error"], "Model output reached the configured generation limit before completion")
        self.assertEqual(record["raw_output"], partial)
        self.assertEqual(record["raw_response"]["response"], partial)

    def test_no_retry_for_truncation(self) -> None:
        config = test_config()
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        post_calls = []

        def capture_post(url, json, timeout):
            post_calls.append(dict(json))
            return FakeResponse(payload={"response": '{"questions": [', "done": True, "done_reason": "length"}, text="ok")

        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", side_effect=capture_post),
        ):
            record = make_record(experiment_id="exp", lecture=lecture, model="tinyllama", config=config, prompt="prompt")
        self.assertEqual(len(post_calls), 1)
        self.assertEqual(record["attempt_count"], 1)
        self.assertFalse(record["retry_performed"])
        self.assertFalse(record["retryable"])
        self.assertEqual(record["retry_reason"], "truncated_output")

    def test_no_retry_for_empty_http_200_response(self) -> None:
        config = test_config()
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        post_calls = []

        def capture_post(url, json, timeout):
            post_calls.append(dict(json))
            return FakeResponse(payload={"response": "", "done": True}, text="ok")

        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", side_effect=capture_post),
        ):
            record = make_record(experiment_id="exp", lecture=lecture, model="tinyllama", config=config, prompt="prompt")
        self.assertEqual(len(post_calls), 1)
        self.assertFalse(record["retry_performed"])
        self.assertFalse(record["retryable"])
        self.assertEqual(record["retry_reason"], "empty_response")

    def test_generation_success_conditions(self) -> None:
        config = test_config()
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        output = {"questions": [valid_mcq(1), valid_mcq(2), valid_mcq(3), valid_mcq(4)]}
        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", return_value=FakeResponse(payload={"response": json.dumps(output), "done": True, "done_reason": "stop", "prompt_eval_count": 100}, text="ok")),
        ):
            record = make_record(experiment_id="exp", lecture=lecture, model="tinyllama", config=config, prompt="prompt")
        self.assertTrue(record["request_success"])
        self.assertTrue(record["generation_completed"])
        self.assertTrue(record["valid_json"])
        self.assertTrue(record["schema_compliant"])
        self.assertTrue(record["complete_question_set"])
        self.assertTrue(record["generation_success"])
        self.assertIsNone(record["schema_failure_reason"])
        self.assertEqual(record["schema_violations"], [])
        self.assertIsNone(record["error"])

    def test_schema_diagnostics_are_recorded_on_result_record(self) -> None:
        config = test_config()
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        output = [valid_mcq(1), valid_mcq(2), valid_mcq(3), valid_mcq(4)]
        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", return_value=FakeResponse(payload={"response": json.dumps(output), "done": True, "done_reason": "stop"}, text="ok")),
        ):
            record = make_record(experiment_id="exp", lecture=lecture, model="tinyllama", config=config, prompt="prompt")
        self.assertTrue(record["valid_json"])
        self.assertFalse(record["schema_compliant"])
        self.assertEqual(record["schema_failure_reason"], "root_not_object")
        self.assertEqual(record["schema_violations"][0]["reason"], "root_not_object")
        self.assertFalse(record["generation_success"])

    def test_cloud_record_uses_model_configuration_without_local_resource_metrics(self) -> None:
        config = EvaluationConfig(
            model_a="gpt-oss-20b-cloud",
            model_b="",
            model_ids=("gpt-oss-20b-cloud",),
            ollama_base_url="http://localhost:11434",
            timeout_seconds=1,
            request_delay_seconds=0,
        )
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        output = {"questions": [valid_mcq(1), valid_mcq(2), valid_mcq(3), valid_mcq(4)]}
        raw_payload = {
            "model": "gpt-oss:20b-cloud",
            "response": json.dumps(output),
            "done": True,
            "done_reason": "stop",
            "prompt_eval_count": 20,
            "eval_count": 40,
            "total_duration": 123,
        }
        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", return_value=FakeResponse(payload=raw_payload, text="ok")),
        ):
            record = make_record(
                experiment_id="exp",
                lecture=lecture,
                model="gpt-oss-20b-cloud",
                config=config,
                prompt="prompt",
                attempt_number=2,
            )
        self.assertEqual(record["model"], "gpt-oss-20b-cloud")
        self.assertEqual(record["exact_model_tag"], "gpt-oss:20b-cloud")
        self.assertEqual(record["deployment_type"], "cloud")
        self.assertTrue(record["internet_required"])
        self.assertTrue(record["data_leaves_institution"])
        self.assertIsNone(record["local_resource_metrics"])
        self.assertEqual(record["cloud_usage"]["eval_count"], 40)
        self.assertEqual(record["pricing_display"]["direct_request_cost"], "Not reported")
        self.assertEqual(record["attempt_number"], 2)
        self.assertIn("request_id", record)
        self.assertIn("prompt_hash", record)
        self.assertTrue(record["format_validation"]["schema_compliant"])

    def test_account_plan_or_usage_limit_errors_are_failed_attempts(self) -> None:
        config = EvaluationConfig(
            model_a="gpt-oss-20b-cloud",
            model_b="",
            model_ids=("gpt-oss-20b-cloud",),
            ollama_base_url="http://localhost:11434",
            timeout_seconds=1,
            request_delay_seconds=0,
        )
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", return_value=FakeResponse(status_code=429, payload={"error": "usage limit reached"}, text="usage limit reached")),
        ):
            record = make_record(
                experiment_id="exp",
                lecture=lecture,
                model="gpt-oss-20b-cloud",
                config=config,
                prompt="prompt",
            )
        self.assertFalse(record["generation_success"])
        self.assertEqual(record["error_category"], "account_or_usage_limit")
        self.assertTrue(record["account_or_usage_limit_error"])
        self.assertFalse(record["retryable"])

    def test_experiment_level_summary_counts(self) -> None:
        records = [
            {"model": "tinyllama", "request_success": True, "empty_response": True, "truncated": False, "valid_json": False, "schema_compliant": False, "complete_question_set": False, "generation_success": False, "retry_performed": False},
            {"model": "tinyllama", "request_success": True, "empty_response": False, "truncated": True, "valid_json": False, "schema_compliant": False, "complete_question_set": False, "generation_success": False, "retry_performed": False},
            {"model": "llama3.1:8b", "request_success": True, "empty_response": False, "truncated": False, "valid_json": True, "schema_compliant": True, "complete_question_set": True, "generation_success": True, "retry_performed": True},
            {"model": "llama3.1:8b", "request_success": False, "empty_response": False, "truncated": False, "valid_json": False, "schema_compliant": False, "complete_question_set": False, "generation_success": False, "retry_performed": False},
        ]
        summary = summarize_records(records, ("tinyllama", "llama3.1:8b"))
        self.assertEqual(summary["overall"]["total_requests"], 4)
        self.assertEqual(summary["overall"]["empty_responses"], 1)
        self.assertEqual(summary["overall"]["truncated_outputs"], 1)
        self.assertEqual(summary["overall"]["generation_successes"], 1)
        self.assertEqual(summary["overall"]["final_http_failures"], 1)
        self.assertEqual(summary["by_model"]["tinyllama"]["total_requests"], 2)

    def test_context_pressure_warning_is_recorded(self) -> None:
        config = test_config(num_ctx=8, num_predict=7)
        lecture = {"id": "lecture_1", "topic": "A", "content": "Content A"}
        output = {"questions": [valid_mcq(1), valid_mcq(2), valid_mcq(3), valid_mcq(4)]}
        with (
            patch("app.llm_evaluation.run_comparison.get_running_models", return_value=([], None)),
            patch("app.llm_evaluation.run_comparison.unload_model"),
            patch("app.llm_evaluation.run_comparison.requests.post", return_value=FakeResponse(payload={"response": json.dumps(output), "done": True, "prompt_eval_count": 4}, text="ok")),
            patch("builtins.print") as print_mock,
        ):
            record = make_record(experiment_id="exp", lecture=lecture, model="tinyllama", config=config, prompt="one two three four five six")
        self.assertTrue(record["context_pressure_warning"])
        self.assertEqual(record["prompt_token_count_source"], "ollama")
        print_mock.assert_any_call("Configured context may be too small: prompt tokens + requested output tokens exceed num_ctx.")

    def test_prompt_size_diagnostics_detect_context_overflow(self) -> None:
        config = test_config(num_ctx=16, num_predict=12)
        lecture = {"id": "lecture_1", "topic": "A", "content": "one two three four five six seven eight nine ten"}
        prompt = "one two three four five six seven eight nine ten"
        diagnostics = prompt_size_diagnostics(lecture, config, prompt)
        self.assertTrue(diagnostics["context_overflow_estimated"])
        self.assertEqual(diagnostics["configured_context_tokens"], 16)
        self.assertEqual(diagnostics["reserved_output_tokens"], 12)
        self.assertFalse(diagnostics["input_chunked"])

    def test_prompt_size_comparison_records_source_and_final_prompt_lengths(self) -> None:
        config = test_config()
        lecture = {"id": "lecture_1", "topic": "A", "content": "word " * 1000}
        prompt = "prompt " * 20
        diagnostics = prompt_size_diagnostics(lecture, config, prompt)
        self.assertGreater(diagnostics["source_input_characters"], diagnostics["normalized_input_characters"])
        self.assertEqual(diagnostics["final_prompt_characters"], len(prompt))
        self.assertEqual(diagnostics["input_handling"], "truncated_to_prompt_chars")

    def test_human_review_excludes_invalid_or_failed_outputs(self) -> None:
        complete_output = {"questions": [valid_mcq(1), valid_mcq(2), valid_mcq(3), valid_mcq(4)]}
        payload = {
            "records": [
                {
                    "lecture_id": "lecture_1",
                    "topic": "A",
                    "model": "tinyllama",
                    "lecture_content": "Content",
                    "valid_json": True,
                    "schema_compliant": False,
                    "generation_success": False,
                    "parsed_output": {"questions": [valid_mcq(1)]},
                },
                {
                    "lecture_id": "lecture_1",
                    "topic": "A",
                    "model": "tinyllama",
                    "lecture_content": "Content",
                    "valid_json": True,
                    "schema_compliant": True,
                    "generation_success": True,
                    "parsed_output": complete_output,
                },
            ]
        }
        rows = build_review_rows(payload, {"Model A": "tinyllama"})
        self.assertEqual(len(rows), 4)
        self.assertTrue(all(row["anonymous_model"] == "Model A" for row in rows))

    def test_human_review_exports_only_successful_aggregates_when_present(self) -> None:
        complete_output = {"questions": [valid_mcq(1), valid_mcq(2), valid_mcq(3), valid_mcq(4)]}
        payload = {
            "records": [
                {**child_question_record(1), "model": "tinyllama", "parsed_output": valid_one_question_output(1)},
                {
                    "record_type": "aggregate_question_set",
                    "aggregate_record": True,
                    "lecture_id": "lecture_1",
                    "topic": "A",
                    "model": "tinyllama",
                    "lecture_content": "Content",
                    "valid_json": True,
                    "schema_compliant": False,
                    "generation_success": False,
                    "parsed_output": {"questions": [valid_mcq(1)]},
                },
                {
                    "record_type": "aggregate_question_set",
                    "aggregate_record": True,
                    "lecture_id": "lecture_1",
                    "topic": "A",
                    "model": "tinyllama",
                    "lecture_content": "Content",
                    "valid_json": True,
                    "schema_compliant": True,
                    "generation_success": True,
                    "parsed_output": complete_output,
                },
            ]
        }
        rows = build_review_rows(payload, {"Model A": "tinyllama"})
        self.assertEqual(len(rows), 4)
        self.assertEqual(sorted(row["question_number"] for row in rows), [1, 2, 3, 4])

    def test_rq1_diagnostic_uses_same_request_code_and_reports_validation(self) -> None:
        config = EvaluationConfig(
            model_a="tinyllama",
            model_b="",
            model_ids=("tinyllama",),
            ollama_base_url="http://localhost:11434",
            timeout_seconds=1,
            request_delay_seconds=0,
        )
        with patch(
            "app.llm_evaluation.diagnose_model.perform_generate_request",
            return_value={"raw_response": {"response": "plain sentence", "done": True, "done_reason": "stop"}, "elapsed_seconds": 0.2, "http_status": 200, "error": None},
        ):
            result = run_rq1_diagnostic(config, "tinyllama", DIAGNOSTIC_PROMPT)
        self.assertEqual(result["endpoint"], config.generate_url)
        self.assertEqual(result["request_payload"], build_generate_payload(config, "tinyllama", DIAGNOSTIC_PROMPT))
        self.assertEqual(result["done"], True)
        self.assertEqual(result["done_reason"], "stop")
        self.assertEqual(result["validation"]["validation_result"], "failed")

    def test_validation_output_categories(self) -> None:
        complete_output = {"questions": [valid_mcq(1), valid_mcq(2), valid_mcq(3), valid_mcq(4)]}
        records = [
            {"model": "tinyllama", "request_success": False, "error_category": "http_failure", "valid_json": False, "parsed_output": None, "empty_response": False, "truncated": False, "schema_compliant": False, "complete_question_set": False, "generation_success": False},
            {"model": "tinyllama", "request_success": True, "error_category": "empty_response", "valid_json": False, "parsed_output": None, "empty_response": True, "truncated": False, "schema_compliant": False, "complete_question_set": False, "generation_success": False},
            {"model": "tinyllama", "request_success": True, "error_category": "truncated_output", "valid_json": False, "parsed_output": None, "empty_response": False, "truncated": True, "schema_compliant": False, "complete_question_set": False, "generation_success": False},
            {"model": "tinyllama", "request_success": True, "error_category": "invalid_json", "valid_json": False, "parsed_output": None, "empty_response": False, "truncated": False, "schema_compliant": False, "complete_question_set": False, "generation_success": False},
            {"model": "tinyllama", "request_success": True, "error_category": "schema_non_compliance", "valid_json": True, "parsed_output": {"questions": [{"type": "essay"}]}, "empty_response": False, "truncated": False, "schema_compliant": False, "complete_question_set": False, "generation_success": False},
            {"model": "tinyllama", "request_success": True, "error_category": "schema_non_compliance", "valid_json": True, "parsed_output": {"questions": [valid_mcq(1)]}, "empty_response": False, "truncated": False, "schema_compliant": False, "complete_question_set": False, "generation_success": False},
            {"model": "tinyllama", "request_success": True, "error_category": None, "valid_json": True, "parsed_output": complete_output, "empty_response": False, "truncated": False, "schema_compliant": True, "complete_question_set": True, "generation_success": True},
        ]
        metrics = analyze_model_records(records)
        self.assertEqual(metrics["request_success_rate"], 85.71)
        self.assertEqual(metrics["empty_response_rate"], 14.29)
        self.assertEqual(metrics["truncation_rate"], 14.29)
        self.assertEqual(metrics["valid_json_rate"], 42.86)
        self.assertEqual(metrics["schema_compliance_rate"], 14.29)
        self.assertEqual(metrics["complete_question_rate"], 14.29)
        self.assertEqual(metrics["generation_success_rate"], 14.29)

    def test_automatic_metrics_do_not_double_count_individual_and_aggregate_records(self) -> None:
        complete_output = {"questions": [valid_mcq(1), valid_mcq(2), valid_mcq(3), valid_mcq(4)]}
        payload = {
            "experiment_id": "exp",
            "records": [
                {**child_question_record(1), "model": "tinyllama"},
                {**child_question_record(2), "model": "tinyllama"},
                {**child_question_record(3), "model": "tinyllama"},
                {**child_question_record(4), "model": "tinyllama"},
                {
                    "record_type": "aggregate_question_set",
                    "aggregate_record": True,
                    "model": "tinyllama",
                    "request_success": True,
                    "valid_json": True,
                    "schema_compliant": True,
                    "complete_question_set": True,
                    "generation_success": True,
                    "parsed_output": complete_output,
                    "empty_response": False,
                    "truncated": False,
                    "elapsed_seconds": 1.0,
                },
            ],
        }
        final_records = final_output_records(payload["records"])
        metrics = calculate_metrics(payload)[0]
        self.assertEqual(len(final_records), 1)
        self.assertEqual(metrics["request_count"], 1)
        self.assertEqual(metrics["individual_request_count"], 4)
        self.assertEqual(metrics["aggregate_request_count"], 1)
        self.assertEqual(metrics["individual_request_success_rate"], 100.0)
        self.assertEqual(metrics["aggregate_generation_success_rate"], 100.0)

    def test_calculation_ignores_missing_optional_ratings(self) -> None:
        df = pd.DataFrame(
            [
                {
                    "anonymous_model": "Model A",
                    "model": "tinyllama",
                    "answer_correct": 1,
                    "grounded_in_lecture": 1,
                    "relevance_score": 4,
                    "clarity_score": 5,
                    "distractor_quality_score": 3,
                    "bloom_label_correct": 1,
                    "difficulty_label_correct": 0,
                },
                {
                    "anonymous_model": "Model A",
                    "model": "tinyllama",
                    "answer_correct": 0,
                    "grounded_in_lecture": 1,
                    "relevance_score": 5,
                    "clarity_score": 4,
                    "distractor_quality_score": "",
                    "bloom_label_correct": 1,
                    "difficulty_label_correct": 1,
                },
            ]
        )
        validate_ranges(df)
        rows = calculate_per_model(df)
        self.assertEqual(rows[0]["average_distractor_quality_score"], 3.0)
        self.assertEqual(rows[0]["reviewed_questions"], 2)

    def test_calibration_pair_parser(self) -> None:
        self.assertEqual(parse_calibration_pairs("1024x280, 2048:512"), [(1024, 280), (2048, 512)])
        with self.assertRaises(ValueError):
            parse_calibration_pairs("2048x0")

    def test_calibration_ranking_recommends_highest_stable_pair(self) -> None:
        rows = [
            {
                "num_ctx": 2048,
                "num_predict": 512,
                "both_models_generation_success": True,
                "generation_successes": 2,
                "http_successes": 2,
                "valid_json_outputs": 2,
                "truncated_outputs": 0,
                "average_tokens_per_second": 8.0,
            },
            {
                "num_ctx": 4096,
                "num_predict": 768,
                "both_models_generation_success": True,
                "generation_successes": 2,
                "http_successes": 2,
                "valid_json_outputs": 2,
                "truncated_outputs": 0,
                "average_tokens_per_second": 4.0,
            },
            {
                "num_ctx": 4096,
                "num_predict": 1024,
                "both_models_generation_success": False,
                "generation_successes": 1,
                "http_successes": 2,
                "valid_json_outputs": 1,
                "truncated_outputs": 1,
                "average_tokens_per_second": 5.0,
            },
        ]
        ranked = rank_rows(rows)
        self.assertEqual(ranked[0]["num_ctx"], 4096)
        self.assertEqual(ranked[0]["num_predict"], 768)
        self.assertTrue(ranked[0]["recommended"])

    def test_calibration_runs_one_lecture_for_both_models_and_writes_ranking(self) -> None:
        config = test_config()
        passages = [
            {"id": "lecture_1", "topic": "A", "content": "Content A"},
            {"id": "lecture_2", "topic": "B", "content": "Content B"},
        ]
        calls: list[tuple[str, int, int, str]] = []

        def fake_record(*, experiment_id, lecture, model, config, prompt, other_models):
            calls.append((model, config.num_ctx, config.num_predict, lecture["id"]))
            success = config.num_ctx == 2048
            return {
                "experiment_id": experiment_id,
                "lecture_id": lecture["id"],
                "model": model,
                "request_success": True,
                "generation_success": success,
                "valid_json": success,
                "truncated": not success,
                "elapsed_seconds": 1.5,
                "tokens_per_second": 9.0 if success else 3.0,
            }

        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "inputs.json"
            results_dir = Path(temp_dir) / "results"
            input_path.write_text(json.dumps(passages), encoding="utf-8")
            with (
                patch("app.llm_evaluation.calibration_benchmark.preflight_checks", return_value={}),
                patch("app.llm_evaluation.calibration_benchmark.unload_and_confirm", return_value={"confirmed": True}),
                patch("app.llm_evaluation.calibration_benchmark.make_record", side_effect=fake_record),
            ):
                ranking_path, detail_path = run_calibration(
                    input_path=input_path,
                    results_dir=results_dir,
                    config=config,
                    pairs=[(1024, 280), (2048, 512)],
                )
            ranking_text = ranking_path.read_text(encoding="utf-8")
            details = json.loads(detail_path.read_text(encoding="utf-8"))

        self.assertEqual(len(calls), 4)
        self.assertEqual({call_item[3] for call_item in calls}, {"lecture_1"})
        self.assertIn("num_ctx,num_predict", ranking_text)
        self.assertEqual(details["recommended_configuration"], {"num_ctx": 2048, "num_predict": 512})


if __name__ == "__main__":
    unittest.main()
