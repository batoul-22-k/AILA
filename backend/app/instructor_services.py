import base64
import io
import json
import logging
import random
import re
import socket
import string
import subprocess
import time
import traceback
import urllib.error
import urllib.request
from dataclasses import dataclass
from datetime import datetime, timezone
from urllib.parse import urlsplit, urlunsplit
from pathlib import Path
from uuid import uuid4

import fitz
import qrcode
from fastapi import UploadFile
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from app.config import get_settings
from app.debug_export import model_names_match, write_llm_debug_export
from app.models import InstructorQuestion, new_id

logger = logging.getLogger(__name__)

LATEST_DEBUG_REQUEST_BY_QUESTION_ID: dict[str, dict] = {}
ALLOWED_OLLAMA_PROVIDERS = {"ollama-local", "ollama-cloud"}
PROVIDER_DEPLOYMENT_TYPES = {"ollama-local": "local", "ollama-cloud": "cloud"}


def new_debug_uuid() -> str:
    return str(uuid4())


def remember_question_debug_request(question_id: str | None, request_metadata: dict | None) -> None:
    if question_id and request_metadata:
        LATEST_DEBUG_REQUEST_BY_QUESTION_ID[question_id] = dict(request_metadata)


def latest_question_debug_request(question_id: str | None) -> dict | None:
    if not question_id:
        return None
    return LATEST_DEBUG_REQUEST_BY_QUESTION_ID.get(question_id)


@dataclass(frozen=True)
class OllamaProviderConfig:
    provider: str
    deployment_type: str
    base_url: str
    generate_url: str
    model: str
    api_key: str | None = None
    endpoint_label: str = ""


def normalize_ollama_base_url(value: str) -> str:
    clean = str(value or "").strip().rstrip("/")
    if not clean:
        return ""
    parsed = urlsplit(clean)
    if parsed.path.rstrip("/") == "/api/generate":
        return urlunsplit((parsed.scheme, parsed.netloc, "", "", "")).rstrip("/")
    return clean


def safe_endpoint_label(provider: str, base_url: str) -> str:
    parsed = urlsplit(str(base_url or ""))
    host = parsed.hostname or ("localhost" if provider == "ollama-local" else "configured-cloud")
    return f"{provider}:{host}"


def redact_secret_text(value: str | None) -> str:
    if value is None:
        return ""
    text = str(value)
    settings = get_settings()
    secrets = [settings.ollama_cloud_api_key]
    for secret in secrets:
        if secret:
            text = text.replace(secret, "[REDACTED]")
    text = re.sub(r"(Authorization:\s*Bearer\s+)[^\s'\"]+", r"\1[REDACTED]", text, flags=re.IGNORECASE)
    text = re.sub(r"(Bearer\s+)[A-Za-z0-9._~+/=-]+", r"\1[REDACTED]", text, flags=re.IGNORECASE)
    return text


def model_response_preview(raw: dict | None, limit: int = 1000) -> str:
    if not isinstance(raw, dict):
        return ""
    response = str(raw.get("response") or "")
    return redact_secret_text(response[:limit].replace("\n", " "))


def categorize_ollama_error(provider_config: OllamaProviderConfig, http_status: int | None, exc: BaseException | None = None) -> str:
    if isinstance(exc, (TimeoutError, socket.timeout)):
        return "timeout"
    if http_status in {401, 403}:
        return "authentication_failed"
    if http_status == 404:
        return "model_not_found"
    if http_status == 429:
        return "rate_limited"
    if http_status is not None and http_status >= 500:
        return "server_error"
    if isinstance(exc, urllib.error.URLError):
        return "local_ollama_not_running" if provider_config.deployment_type == "local" else "endpoint_unreachable"
    if http_status is not None:
        return "http_error"
    return "request_failed"


def resolve_ollama_provider(provider: str | None = None, model: str | None = None) -> OllamaProviderConfig:
    settings = get_settings()
    selected_provider = (provider or settings.ollama_provider or "ollama-local").strip()
    if selected_provider not in ALLOWED_OLLAMA_PROVIDERS:
        raise ValueError(f"Unsupported Ollama provider: {selected_provider}")

    if selected_provider == "ollama-local":
        base_url = normalize_ollama_base_url(settings.ollama_local_base_url or settings.ollama_url)
        configured_model = settings.ollama_local_model or settings.ollama_model
        api_key = None
    else:
        base_url = normalize_ollama_base_url(settings.ollama_cloud_base_url)
        configured_model = settings.ollama_cloud_model
        api_key = settings.ollama_cloud_api_key or None
        if not base_url:
            raise ValueError("OLLAMA_CLOUD_BASE_URL is not configured.")
        if not configured_model:
            raise ValueError("OLLAMA_CLOUD_MODEL is not configured.")

    requested_model = (model or configured_model or "").strip()
    if not requested_model:
        raise ValueError(f"No model is configured for {selected_provider}.")
    if requested_model != configured_model:
        raise ValueError(f"Model {requested_model!r} is not allowed for provider {selected_provider}.")

    deployment_type = PROVIDER_DEPLOYMENT_TYPES[selected_provider]
    return OllamaProviderConfig(
        provider=selected_provider,
        deployment_type=deployment_type,
        base_url=base_url,
        generate_url=f"{base_url}/api/generate",
        model=configured_model,
        api_key=api_key,
        endpoint_label=safe_endpoint_label(selected_provider, base_url),
    )


def get_ollama_base_url(provider: str | None = None) -> str:
    return resolve_ollama_provider(provider).base_url


def get_ollama_tags_url(provider: str | None = None) -> str:
    return f"{get_ollama_base_url(provider)}/api/tags"


def auth_headers_for_provider(provider_config: OllamaProviderConfig) -> dict[str, str]:
    if provider_config.api_key:
        return {"Authorization": f"Bearer {provider_config.api_key}"}
    return {}


def get_installed_ollama_models(provider: str | None = None) -> list[str]:
    provider_config = resolve_ollama_provider(provider)
    request = urllib.request.Request(
        f"{provider_config.base_url}/api/tags",
        headers=auth_headers_for_provider(provider_config),
        method="GET",
    )
    with urllib.request.urlopen(request, timeout=5) as response:
        raw = json.loads(response.read().decode("utf-8"))
    return [model.get("name", "") for model in raw.get("models", []) if model.get("name")]


def get_ollama_status(provider: str | None = None) -> dict:
    try:
        provider_config = resolve_ollama_provider(provider)
        models = get_installed_ollama_models(provider_config.provider)
        requested = provider_config.model
        has_model = requested in models or f"{requested}:latest" in models
        return {
            "running": True,
            "provider": provider_config.provider,
            "deployment_type": provider_config.deployment_type,
            "endpoint_label": provider_config.endpoint_label,
            "model": requested,
            "models": models,
            "model_available": has_model,
            "message": "Ollama is reachable." if has_model else f"Ollama is running, but model '{requested}' is not pulled.",
        }
    except Exception as exc:
        selected_provider = provider or get_settings().ollama_provider
        model = ""
        endpoint_label = selected_provider or "ollama-local"
        try:
            provider_config = resolve_ollama_provider(provider)
            model = provider_config.model
            endpoint_label = provider_config.endpoint_label
        except Exception:
            model = get_settings().ollama_model
        return {
            "running": False,
            "provider": selected_provider,
            "deployment_type": PROVIDER_DEPLOYMENT_TYPES.get(selected_provider, "local"),
            "endpoint_label": endpoint_label,
            "model": model,
            "models": [],
            "model_available": False,
            "message": f"Ollama is not reachable: {redact_secret_text(str(exc))}",
        }


def start_ollama_server() -> dict:
    status = get_ollama_status("ollama-local")
    if status["running"]:
        return status

    creationflags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    try:
        subprocess.Popen(
            ["ollama", "serve"],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            stdin=subprocess.DEVNULL,
            creationflags=creationflags,
        )
    except FileNotFoundError as exc:
        raise RuntimeError("Ollama is not installed or is not available on PATH.") from exc
    except Exception as exc:
        raise RuntimeError(f"Could not start Ollama: {exc}") from exc

    for _ in range(15):
        time.sleep(1)
        status = get_ollama_status("ollama-local")
        if status["running"]:
            return status

    raise RuntimeError("Started Ollama, but it did not become reachable within 15 seconds.")


def build_ollama_payload(model: str, prompt: str, options: dict, keep_alive: str = "2m") -> dict:
    return {
        "model": model,
        "prompt": prompt,
        "stream": False,
        "format": "json",
        "keep_alive": keep_alive,
        "options": options,
    }


def safe_ollama_request_payload(payload: dict, provider_config: OllamaProviderConfig) -> dict:
    return {
        **payload,
        "provider": provider_config.provider,
        "deployment_type": provider_config.deployment_type,
        "endpoint_label": provider_config.endpoint_label,
    }


def perform_ollama_generate_request(
    *,
    provider: str,
    base_url: str,
    model: str,
    api_key: str | None,
    prompt: str,
    options: dict,
    timeout: int,
    endpoint_label: str | None = None,
) -> dict:
    provider_config = OllamaProviderConfig(
        provider=provider,
        deployment_type=PROVIDER_DEPLOYMENT_TYPES.get(provider, "cloud" if "cloud" in provider else "local"),
        base_url=normalize_ollama_base_url(base_url),
        generate_url=f"{normalize_ollama_base_url(base_url)}/api/generate",
        model=model,
        api_key=api_key,
        endpoint_label=endpoint_label or safe_endpoint_label(provider, base_url),
    )
    payload = build_ollama_payload(model, prompt, options)
    headers = {"Content-Type": "application/json", **auth_headers_for_provider(provider_config)}
    request = urllib.request.Request(
        provider_config.generate_url,
        data=json.dumps(payload).encode("utf-8"),
        headers=headers,
        method="POST",
    )
    request_started = time.perf_counter()
    raw_response_text: str | None = None
    http_status: int | None = None
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            http_status = getattr(response, "status", None)
            raw_response_text = response.read().decode("utf-8")
        elapsed_ms = round((time.perf_counter() - request_started) * 1000)
        parsed_raw = json.loads(raw_response_text)
        if not isinstance(parsed_raw, dict):
            raise RuntimeError("Ollama returned a non-object JSON response.")
        return {
            "provider_config": provider_config,
            "payload": payload,
            "safe_payload": safe_ollama_request_payload(payload, provider_config),
            "raw_response_text": raw_response_text,
            "parsed_response": parsed_raw,
            "http_status": http_status,
            "elapsed_ms": elapsed_ms,
            "error": None,
            "error_category": None,
        }
    except (TimeoutError, socket.timeout) as exc:
        elapsed_ms = round((time.perf_counter() - request_started) * 1000)
        category = categorize_ollama_error(provider_config, None, exc)
        return {
            "provider_config": provider_config,
            "payload": payload,
            "safe_payload": safe_ollama_request_payload(payload, provider_config),
            "raw_response_text": raw_response_text,
            "parsed_response": None,
            "http_status": None,
            "elapsed_ms": elapsed_ms,
            "error": redact_secret_text(str(exc) or "Ollama request timed out."),
            "error_category": category,
            "exception": exc,
        }
    except urllib.error.HTTPError as exc:
        elapsed_ms = round((time.perf_counter() - request_started) * 1000)
        try:
            detail = exc.read().decode("utf-8")
        except Exception:
            detail = str(exc)
        raw_response_text = detail
        category = categorize_ollama_error(provider_config, exc.code, exc)
        return {
            "provider_config": provider_config,
            "payload": payload,
            "safe_payload": safe_ollama_request_payload(payload, provider_config),
            "raw_response_text": raw_response_text,
            "parsed_response": None,
            "http_status": exc.code,
            "elapsed_ms": elapsed_ms,
            "error": redact_secret_text(f"Ollama returned HTTP {exc.code}: {detail}"),
            "error_category": category,
            "exception": exc,
        }
    except urllib.error.URLError as exc:
        elapsed_ms = round((time.perf_counter() - request_started) * 1000)
        category = categorize_ollama_error(provider_config, None, exc)
        return {
            "provider_config": provider_config,
            "payload": payload,
            "safe_payload": safe_ollama_request_payload(payload, provider_config),
            "raw_response_text": raw_response_text,
            "parsed_response": None,
            "http_status": None,
            "elapsed_ms": elapsed_ms,
            "error": redact_secret_text(str(exc)),
            "error_category": category,
            "exception": exc,
        }
    except json.JSONDecodeError as exc:
        elapsed_ms = round((time.perf_counter() - request_started) * 1000)
        return {
            "provider_config": provider_config,
            "payload": payload,
            "safe_payload": safe_ollama_request_payload(payload, provider_config),
            "raw_response_text": raw_response_text,
            "parsed_response": None,
            "http_status": http_status,
            "elapsed_ms": elapsed_ms,
            "error": "Ollama returned malformed JSON.",
            "error_category": "malformed_json",
            "exception": exc,
        }
    except RuntimeError as exc:
        elapsed_ms = round((time.perf_counter() - request_started) * 1000)
        return {
            "provider_config": provider_config,
            "payload": payload,
            "safe_payload": safe_ollama_request_payload(payload, provider_config),
            "raw_response_text": raw_response_text,
            "parsed_response": None,
            "http_status": http_status,
            "elapsed_ms": elapsed_ms,
            "error": redact_secret_text(str(exc)),
            "error_category": "malformed_json",
            "exception": exc,
        }


def provider_config_metadata(provider_config: OllamaProviderConfig, parsed_response: dict | None = None, http_status: int | None = None, elapsed_ms: int | None = None) -> dict:
    response_model = parsed_response.get("model") if isinstance(parsed_response, dict) else None
    return {
        "provider": provider_config.provider,
        "deployment_type": provider_config.deployment_type,
        "requested_model": provider_config.model,
        "response_model": response_model,
        "model_name_match": model_names_match(provider_config.model, response_model) if response_model else None,
        "endpoint_label": provider_config.endpoint_label,
        "processing_time_ms": elapsed_ms,
        "prompt_tokens": parsed_response.get("prompt_eval_count") if isinstance(parsed_response, dict) else None,
        "completion_tokens": parsed_response.get("eval_count") if isinstance(parsed_response, dict) else None,
        "total_tokens": (
            parsed_response.get("prompt_eval_count") + parsed_response.get("eval_count")
            if isinstance(parsed_response, dict)
            and isinstance(parsed_response.get("prompt_eval_count"), int)
            and isinstance(parsed_response.get("eval_count"), int)
            else None
        ),
        "done": parsed_response.get("done") if isinstance(parsed_response, dict) else None,
        "done_reason": parsed_response.get("done_reason") if isinstance(parsed_response, dict) else None,
        "http_status": http_status,
    }


def generated_question_metadata(request_metadata: dict | None, provider_config: OllamaProviderConfig) -> dict:
    metadata = request_metadata or {}
    keys = [
        "provider",
        "deployment_type",
        "requested_model",
        "response_model",
        "model_name_match",
        "endpoint_label",
        "processing_time_ms",
        "prompt_tokens",
        "completion_tokens",
        "total_tokens",
        "done",
        "done_reason",
        "http_status",
    ]
    generated = {key: metadata.get(key) for key in keys}
    generated["provider"] = generated.get("provider") or provider_config.provider
    generated["deployment_type"] = generated.get("deployment_type") or provider_config.deployment_type
    generated["requested_model"] = generated.get("requested_model") or provider_config.model
    generated["endpoint_label"] = generated.get("endpoint_label") or provider_config.endpoint_label
    return generated


def test_ollama_connection(provider: str | None = None, model: str | None = None) -> dict:
    provider_config = resolve_ollama_provider(provider, model)
    settings = get_settings()
    result = perform_ollama_generate_request(
        provider=provider_config.provider,
        base_url=provider_config.base_url,
        model=provider_config.model,
        api_key=provider_config.api_key,
        prompt='Return {"ok": true}.',
        options={
            "temperature": 0,
            "num_ctx": min(settings.ollama_num_ctx, 1024),
            "num_predict": 16,
            "num_gpu": settings.ollama_num_gpu,
        },
        timeout=min(settings.ollama_timeout_seconds, 30),
        endpoint_label=provider_config.endpoint_label,
    )
    parsed = result.get("parsed_response")
    reachable = result.get("http_status") == 200 and isinstance(parsed, dict) and parsed.get("done") is True
    error_category = result.get("error_category")
    if result.get("http_status") == 200 and isinstance(parsed, dict):
        if parsed.get("done") is False:
            error_category = "incomplete_response"
        elif parsed.get("done") is True:
            error_category = None
    return {
        "provider": provider_config.provider,
        "configured_model": provider_config.model,
        "deployment_type": provider_config.deployment_type,
        "endpoint_label": provider_config.endpoint_label,
        "reachable": reachable,
        "http_status": result.get("http_status"),
        "elapsed_time_ms": result.get("elapsed_ms"),
        "elapsed_seconds": round((result.get("elapsed_ms") or 0) / 1000, 4),
        "response_model": parsed.get("model") if isinstance(parsed, dict) else None,
        "error_category": error_category,
    }


def get_storage_root() -> Path:
    root = Path(get_settings().storage_dir)
    root.mkdir(parents=True, exist_ok=True)
    (root / "uploads").mkdir(exist_ok=True)
    (root / "presentations").mkdir(exist_ok=True)
    return root


async def save_upload_file(file: UploadFile, upload_id: str) -> Path:
    suffix = Path(file.filename or "lecture").suffix.lower()
    target = get_storage_root() / "uploads" / f"{upload_id}{suffix}"
    target.write_bytes(await file.read())
    return target


def extract_readable_content(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
      return extract_pdf_text(path)
    if suffix == ".pptx":
      return extract_pptx_text(path)
    raise ValueError("Only PPTX and PDF files are supported.")


def extract_pdf_text(path: Path) -> str:
    lines: list[str] = []
    with fitz.open(path) as document:
        for page_index, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if text:
                lines.append(f"[Page {page_index}]\n{text}")
    return "\n\n".join(lines)


def extract_pptx_text(path: Path) -> str:
    presentation = Presentation(path)
    lines: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        slide_lines: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_lines.append(shape.text.strip())
        if slide_lines:
            lines.append(f"[Slide {slide_index}]\n" + "\n".join(slide_lines))
    return "\n\n".join(lines)


def clean_extracted_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = text.replace("\x0b", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


GENERIC_LECTURE_TERMS = {
    "about",
    "answer",
    "branch",
    "content",
    "course",
    "engineering",
    "faculty",
    "lecture",
    "material",
    "page",
    "prepared",
    "project",
    "question",
    "slide",
    "student",
    "supervised",
    "thank",
    "university",
}

SHORT_ANSWER_MAX_WORDS = 14
QUESTION_TEXT_MAX_WORDS = 25
OPTION_MAX_WORDS = 15
QUESTION_LEAKAGE_STOP_WORDS = GENERIC_LECTURE_TERMS | {
    "apply",
    "could",
    "describe",
    "does",
    "explain",
    "example",
    "guide",
    "helps",
    "identify",
    "important",
    "instructor",
    "mean",
    "next",
    "sentence",
    "short",
    "words",
}
GENERIC_AI_QUESTION_PATTERNS = (
    r"\bwhich\s+option\s+best\s+explains\s+the\s+impact\s+of\s+this\s+slide",
    r"\bwhich\s+choice\s+best\s+applies\s+the\s+lecture'?s\s+idea\b",
    r"\bwhat\s+action\s+or\s+conclusion\s+follows\s+from\s+the\s+lecture\s+concept\b",
    r"\bwhat\s+key\s+conclusion\s+follows\s+from\s+the\s+slide\s+about\b",
    r"\ba\s+complete\s+question\s+based\s+on\s+the\s+lecture\s+content\b",
)


def parse_slide_sections(text: str) -> list[dict]:
    matches = list(re.finditer(r"\[Slide\s+(\d+)\]\s*", text, flags=re.IGNORECASE))
    sections: list[dict] = []
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        raw_lines = text[start:end].replace("\x0b", "\n").splitlines()
        lines = [re.sub(r"\s+", " ", line).strip(" -:\u2022") for line in raw_lines]
        lines = [line for line in lines if line]
        if not lines:
            continue
        title_index = 0
        for line_index, line in enumerate(lines):
            if re.search(r"\b(requirements?|project needs|technical approach|budget|schedule|selection criteria)\b", line, flags=re.IGNORECASE):
                title_index = line_index
                break
        title = lines[title_index]
        body_lines = lines[:title_index] + lines[title_index + 1 :]
        body = [line for line in body_lines if not is_low_value_slide_line(line)]
        sections.append({"slide": int(match.group(1)), "title": title, "body": body, "lines": lines})
    return sections


def is_low_value_slide_line(line: str) -> bool:
    normalized = re.sub(r"[^a-z0-9]+", " ", line.lower()).strip()
    if not normalized:
        return True
    if normalized.isdigit():
        return True
    if len(normalized) <= 2:
        return True
    if normalized in GENERIC_LECTURE_TERMS:
        return True
    if any(term in normalized for term in ("prepared by", "supervised by", "faculty", "university", "branch iii")):
        return True
    if re.search(r"\b\d{3,}\b", normalized) and not any(term in normalized for term in ("budget", "price", "date", "oct", "nov", "dec")):
        return True
    if re.fullmatch(r"[0-9\s%$.,]+", line):
        return True
    return False


def is_low_value_slide_section(section: dict) -> bool:
    title = str(section.get("title", "")).lower()
    body = section.get("body") or []
    low_value_titles = ("thank", "questions", "sneak peek")
    metadata_terms = ("prepared by", "supervised by", "faculty", "university")
    if any(term in title for term in low_value_titles):
        return True
    if any(term in " ".join(section.get("lines") or []).lower() for term in metadata_terms) and len(body) < 2:
        return True
    return len(" ".join(body)) < 35


def summarize_slide_body(body: list[str], limit: int = 220) -> str:
    summary = " ".join(body)
    summary = re.sub(r"\s+", " ", summary).strip()
    if len(summary) <= limit:
        return summary
    return summary[:limit].rsplit(" ", 1)[0].rstrip(".,;:") + "."


def meaningful_slide_sections(text: str) -> list[dict]:
    sections = [section for section in parse_slide_sections(text) if not is_low_value_slide_section(section)]
    return sections or parse_slide_sections(text)


def select_question_section(text: str, question_index: int | None, avoid_questions: list[str] | None = None) -> dict | None:
    sections = meaningful_slide_sections(text)
    if not sections:
        return None
    avoid_text = " ".join(avoid_questions or []).lower()
    offset = max((question_index or 1) - 1, 0)
    for step in range(len(sections)):
        section = sections[(offset + step) % len(sections)]
        title = section["title"].lower()
        body = summarize_slide_body(section.get("body") or []).lower()
        if title not in avoid_text and body not in avoid_text:
            return section
    return sections[offset % len(sections)]


def build_prompt_context(text: str, limit: int) -> str:
    sections = meaningful_slide_sections(text)
    if not sections:
        return text[:limit]
    chunks: list[str] = []
    for section in sections:
        body = summarize_slide_body(section.get("body") or section.get("lines") or [], limit=360)
        if not body:
            continue
        chunks.append(f"[Slide {section['slide']}] {section['title']}\n{body}")
        if len("\n\n".join(chunks)) >= limit:
            break
    return "\n\n".join(chunks)[:limit]


WEAK_QUESTION_PATTERNS = (
    r"\bwhat\s+is\s+the\s+lecture\s+content\b",
    r"\bwhat\s+(?:are|is)\s+we\s+going\s+to\s+talk\s+about\b",
    r"\bwhat\s+is\s+this\s+lecture\s+about\b",
    r"\bwhat\s+is\s+the\s+main\s+topic\s+of\s+the\s+lecture\b",
    r"\ba\s+complete\s+question\s+based\s+on\s+the\s+lecture\s+content\b",
)


def is_weak_question_text(question_text: str, avoid_questions: list[str] | None = None) -> bool:
    normalized = re.sub(r"\s+", " ", str(question_text or "").strip().lower())
    if not normalized:
        return True
    if any(re.search(pattern, normalized) for pattern in WEAK_QUESTION_PATTERNS):
        return True
    previous_questions = {str(question).strip().lower() for question in avoid_questions or []}
    if normalized in previous_questions:
        return True
    words = re.findall(r"[a-z0-9]+", normalized)
    content_words = [word for word in words if word not in GENERIC_LECTURE_TERMS and len(word) > 3]
    content_set = set(content_words)
    for previous in previous_questions:
        previous_words = {
            word
            for word in re.findall(r"[a-z0-9]+", previous)
            if word not in GENERIC_LECTURE_TERMS and len(word) > 3
        }
        if not previous_words or not content_set:
            continue
        overlap = len(content_set & previous_words) / max(len(content_set | previous_words), 1)
        if overlap >= 0.65:
            return True
    return len(content_words) < 3


def concise_short_answer(text: str, max_words: int = SHORT_ANSWER_MAX_WORDS) -> str:
    cleaned = re.sub(r"\s+", " ", str(text or "")).strip(" .")
    if not cleaned:
        return ""
    first_sentence = re.split(r"(?<=[.!?])\s+", cleaned)[0].strip(" .")
    first_clause = re.split(r"\s+(?:because|since|as|which|where|when|while)\s+", first_sentence, maxsplit=1, flags=re.IGNORECASE)[0]
    words = first_clause.split()
    if len(words) <= max_words:
        return first_clause.strip(" .")
    return " ".join(words[:max_words]).strip(" .,;:")


def word_count(text: str) -> int:
    return len(re.findall(r"\b[\w'-]+\b", str(text or "")))


def truncate_words(text: str, max_words: int) -> str:
    words = re.findall(r"\S+", str(text or "").strip())
    if len(words) <= max_words:
        return str(text or "").strip()
    return " ".join(words[:max_words]).strip(" .,;:") + "."


def title_case_concept(text: str) -> str:
    cleaned = re.sub(r"\s+", " ", str(text or "").strip(" .:-"))
    cleaned = re.sub(r"^(slide|topic|concept)\s+\d+\s*[:.-]?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\b(example|examples)\b.*$", "", cleaned, flags=re.IGNORECASE).strip(" .:-")
    return truncate_words(cleaned or "this concept", 8).rstrip(".")


def has_multiple_sentences(text: str) -> bool:
    sentences = [part for part in re.split(r"[.!?]+", str(text or "")) if part.strip()]
    return len(sentences) > 1


def is_generic_question_text(question_text: str) -> bool:
    normalized = re.sub(r"\s+", " ", str(question_text or "").strip().lower())
    return any(re.search(pattern, normalized) for pattern in GENERIC_AI_QUESTION_PATTERNS)


def slide_sentences(text: str) -> list[str]:
    cleaned = re.sub(r"\[[^\]]+\]", " ", str(text or ""))
    candidates = re.split(r"(?<=[.!?])\s+|\n+", cleaned)
    return [
        re.sub(r"\s+", " ", candidate).strip(" .:-")
        for candidate in candidates
        if word_count(candidate) >= 8
    ]


def near_long_slide_text(value: str, lecture_text: str) -> bool:
    normalized_value = re.sub(r"\s+", " ", str(value or "").strip().lower())
    if word_count(normalized_value) < 8:
        return False
    value_terms = set(content_terms(normalized_value))
    if len(value_terms) < 5:
        return False
    for sentence in slide_sentences(lecture_text):
        normalized_sentence = sentence.lower()
        if normalized_value and normalized_value in normalized_sentence:
            return True
        sentence_terms = set(content_terms(normalized_sentence))
        if len(sentence_terms) >= 5 and len(value_terms & sentence_terms) / max(len(value_terms), 1) >= 0.85:
            return True
    return False


def content_terms(text: str) -> list[str]:
    return [
        term
        for term in re.findall(r"[a-z0-9]+", str(text or "").lower())
        if term not in QUESTION_LEAKAGE_STOP_WORDS and len(term) > 2
    ]


def answer_leaks_into_question(question_text: str, correct_answer: str) -> bool:
    question = re.sub(r"\s+", " ", str(question_text or "").strip().lower())
    answer = re.sub(r"\s+", " ", str(correct_answer or "").strip(" .").lower())
    if not question or not answer:
        return False

    answer_words = answer.split()
    if len(answer) >= 18 and answer in question:
        return True
    if len(answer_words) >= 4:
        answer_phrases = [" ".join(answer_words[index : index + 4]) for index in range(len(answer_words) - 3)]
        if any(phrase in question for phrase in answer_phrases):
            return True

    answer_terms = set(content_terms(answer))
    question_terms = set(content_terms(question))
    if len(answer_terms) >= 3:
        overlap = len(answer_terms & question_terms) / max(len(answer_terms), 1)
        return overlap >= 0.75
    return False


def concise_option(text: str, fallback: str) -> str:
    cleaned = re.sub(r"\s+", " ", str(text or "").strip())
    cleaned = re.sub(r"\b(example|for example)\b.*$", "", cleaned, flags=re.IGNORECASE).strip(" .:-")
    if not cleaned:
        cleaned = fallback
    first_sentence = re.split(r"[.!?]+", cleaned, maxsplit=1)[0].strip(" .:-")
    return truncate_words(first_sentence or fallback, OPTION_MAX_WORDS).rstrip(".")


def clean_correct_answer(text: str, concept: str) -> str:
    cleaned = concise_short_answer(text, max_words=OPTION_MAX_WORDS)
    if not cleaned or cleaned.lower() in {"brief expected answer", "definition"}:
        return f"{concept} describes the main idea"
    return cleaned


def misconception_options(concept: str) -> list[str]:
    return [
        f"{concept} depends on storage order",
        f"{concept} changes when examples are reordered",
        f"{concept} is defined by one sample value",
    ]


def shuffled_mcq_options(options: list[str], correct_answer: str) -> list[str]:
    normalized_correct = str(correct_answer or "").strip()
    clean_options: list[str] = []
    for option in options:
        clean_option = str(option or "").strip()
        if clean_option and clean_option.lower() not in {existing.lower() for existing in clean_options}:
            clean_options.append(clean_option)
    if normalized_correct and normalized_correct.lower() not in {option.lower() for option in clean_options}:
        clean_options.insert(0, normalized_correct)
    if normalized_correct:
        protected = [option for option in clean_options if option.lower() == normalized_correct.lower()]
        remaining = [option for option in clean_options if option.lower() != normalized_correct.lower()]
        clean_options = [*protected[:1], *remaining[:3]]
    else:
        clean_options = clean_options[:4]
    shuffled = clean_options[:]
    random.shuffle(shuffled)
    return shuffled


def local_clean_question_from_lecture(
    text: str,
    question_type: str | None,
    bloom_level: str | None,
    difficulty: str | None,
    question_index: int | None = None,
    avoid_questions: list[str] | None = None,
) -> dict:
    section = select_question_section(text, question_index, avoid_questions)
    normalized_type = question_type or "mcq"
    title = title_case_concept(section["title"] if section else "")
    body = section.get("body") or section.get("lines") or [] if section else []
    body_summary = summarize_slide_body(body, limit=120) if body else ""
    concept = title_case_concept(title or body_summary or "this concept")
    correct = clean_correct_answer(body_summary or title, concept)
    source_slide = section["slide"] if section else None

    if normalized_type == "short_answer":
        return {
            "type": "short_answer",
            "question_text": truncate_words(f"Explain why {concept} matters in this lesson.", QUESTION_TEXT_MAX_WORDS),
            "options": [],
            "correct_answer": concise_short_answer(correct),
            "explanation": f"{concept} is the key idea from the selected slide.",
            "bloom_level": bloom_level or "Understand",
            "difficulty": difficulty or "Medium",
            "source_slide": source_slide,
        }

    correct_option = concise_option(correct, f"{concept} describes the main idea")
    distractors = [concise_option(option, option) for option in misconception_options(concept)]
    options = shuffled_mcq_options([correct_option, *distractors[:3]], correct_option)
    return {
        "type": "mcq",
        "question_text": truncate_words(f"Which statement correctly describes {concept}?", QUESTION_TEXT_MAX_WORDS),
        "options": options,
        "correct_answer": correct_option,
        "explanation": f"{correct_option} is the concise statement supported by the lecture.",
        "bloom_level": bloom_level or "Understand",
        "difficulty": difficulty or "Medium",
        "source_slide": source_slide,
    }


def question_quality_errors(question: dict, lecture_text: str) -> list[str]:
    errors: list[str] = []
    question_type = question.get("type", "mcq")
    question_text = str(question.get("question_text") or "")
    if word_count(question_text) > QUESTION_TEXT_MAX_WORDS:
        errors.append("question_too_long")
    if is_generic_question_text(question_text):
        errors.append("generic_question")
    if question_type == "short_answer" and answer_leaks_into_question(question_text, question.get("correct_answer", "")):
        errors.append("answer_leak")

    if question_type == "mcq":
        options = [str(option or "").strip() for option in question.get("options") or []]
        if len(options) != 4:
            errors.append("wrong_option_count")
        if question.get("correct_answer") not in options:
            errors.append("missing_correct_option")
        option_lengths = [word_count(option) for option in options]
        if any(length > OPTION_MAX_WORDS for length in option_lengths):
            errors.append("option_too_long")
        if any(has_multiple_sentences(option) for option in options):
            errors.append("multi_sentence_option")
        if any(near_long_slide_text(option, lecture_text) for option in options):
            errors.append("option_copies_slide")
        if len({option.lower() for option in options}) != len(options):
            errors.append("duplicate_options")
        if option_lengths and max(option_lengths) - min(option_lengths) > 9:
            errors.append("unbalanced_options")

    return errors


def enforce_question_quality(
    question: dict,
    lecture_text: str,
    fallback: dict,
) -> tuple[dict, list[str]]:
    repaired = dict(question)
    question_type = repaired.get("type", "mcq")
    repaired["question_text"] = truncate_words(repaired.get("question_text") or fallback["question_text"], QUESTION_TEXT_MAX_WORDS)

    if is_generic_question_text(repaired["question_text"]):
        repaired["question_text"] = fallback["question_text"]

    if question_type == "mcq":
        concept = title_case_concept(fallback["question_text"].replace("Which statement correctly describes", "").strip(" ?"))
        raw_options = [option for option in repaired.get("options", []) if str(option).strip()]
        clean_options = [concise_option(option, "") for option in raw_options]
        clean_options = [
            option
            for option in clean_options
            if option and not has_multiple_sentences(option) and not near_long_slide_text(option, lecture_text)
        ]
        correct = concise_option(repaired.get("correct_answer"), fallback.get("correct_answer", ""))
        if correct and correct not in clean_options and not near_long_slide_text(correct, lecture_text):
            clean_options.insert(0, correct)
        clean_options.extend(fallback.get("options") or misconception_options(concept))

        deduped: list[str] = []
        for option in clean_options:
            if option.lower() not in {existing.lower() for existing in deduped}:
                deduped.append(option)
            if len(deduped) == 4:
                break
        repaired["correct_answer"] = correct if correct in deduped else fallback["correct_answer"]
        repaired["options"] = shuffled_mcq_options(deduped, repaired["correct_answer"])
        if len(repaired["options"]) < 4:
            repaired = fallback
    else:
        repaired["options"] = []
        repaired["correct_answer"] = concise_short_answer(repaired.get("correct_answer") or fallback["correct_answer"])
        if is_generic_question_text(repaired["question_text"]) or answer_leaks_into_question(repaired["question_text"], repaired["correct_answer"]):
            repaired["question_text"] = fallback["question_text"]

    errors = question_quality_errors(repaired, lecture_text)
    return repaired, errors


def build_question_prompt(
    text: str,
    regenerate_question: InstructorQuestion | None = None,
    question_type: str | None = None,
    bloom_level: str | None = None,
    difficulty: str | None = None,
    output_language: str | None = None,
    question_index: int | None = None,
    avoid_questions: list[str] | None = None,
) -> str:
    settings = get_settings()
    base = build_prompt_context(text, settings.ollama_prompt_chars)
    regenerate = ""
    if regenerate_question:
        regenerate = (
            "\nRegenerate only this question while preserving the same lecture topic and improving quality:\n"
            f"{regenerate_question.model_dump_json()}\n"
        )
    avoid = "\n".join(f"- {question}" for question in (avoid_questions or []) if question)
    avoid_block = f"\nDo not repeat or closely paraphrase these previous questions:\n{avoid}\n" if avoid else ""

    if question_type:
        options_instruction = (
            'Use exactly 4 non-empty answer options. Example options: ["Concept A", "Concept B", "Concept C", "Concept D"].'
            if question_type == "mcq"
            else "Use an empty options array: []."
        )
        example_options = '["Definition", "Example", "Process", "Outcome"]' if question_type == "mcq" else "[]"
        example_answer = "Definition" if question_type == "mcq" else "Brief expected answer."
        answer_rule = (
            "correct_answer must match one option"
            if question_type == "mcq"
            else "correct_answer must be only a few words or one short sentence, maximum 14 words"
        )
        return f"""
You are generating one instructor-reviewed classroom engagement question.
Return only minified valid JSON. Do not include markdown fences, commentary, or extra text.
Return exactly one question object inside the questions array.
Do not create a second question.
Do not leave any field empty.
Do not copy the example wording from the JSON shape.
Choose a different lecture concept than any previous question.
Generate classroom engagement questions, not exam-style copy questions.
Use the lecture content as source, but rewrite in instructor-friendly wording.
Keep wording concise and use simple academic language.
Use only concrete facts from the provided slide excerpts.
Ignore cover slides, presenter names, university/faculty metadata, page numbers, and thank-you slides.
Do not ask vague questions such as "What are we going to talk about?"
Never ask "What is the lecture content?", "What is this lecture about?", or any broad summary question.
Ask about a specific decision, consequence, comparison, problem, solution, tradeoff, or scenario from a slide.
Avoid simple definition questions.
Avoid questions based only on slide titles; use the slide details and relationships between ideas.
Prefer understanding and application questions over recall-only questions.
Question text must be 16 words or fewer.
For multiple choice questions, use exactly 4 answer options.
For multiple choice questions, every option must be 6 words or fewer.
For multiple choice questions, options must be short statements, not paragraphs.
For multiple choice questions, options must not contain commas, semicolons, or multiple clauses.
For multiple choice questions, do not paste full slide sentences into options.
For multiple choice questions, keep all answer choices balanced in length.
For multiple choice questions, make every distractor plausible, related to the same concept, and distinct from the correct answer.
For multiple choice questions, do not use random excerpts from other slides as distractors.
For multiple choice questions, avoid obvious wrong answers.
For multiple choice questions, question_text should be direct, such as "Which statement correctly describes a relation schema?"
Do not use generic wording like "Which option best explains the impact of this slide's idea?"
For short answer questions, ask one clear idea in 25 words or fewer.
For short answer questions, prefer direct questions such as "Why is tuple order irrelevant in a relation?"
For short answer questions, do not require a paragraph, long explanation, or multi-step essay response.
For short answer questions, never include the correct answer, answer phrase, or a near-copy of it in question_text.
For short answer questions, do not write "Example:" in question_text or copy example text from the slide into question_text.
Make question_text ask what, why, or how; keep the expected answer hidden in correct_answer only.
Explanation must be 8 words or fewer.
Include source_slide when a slide number is available.
Rule: {answer_rule}.
{options_instruction}

Lecture content:
{base}
{regenerate}
{avoid_block}

Create exactly 1 {question_type} question.
Question number in this set: {question_index or 1}
Bloom level: {bloom_level or "Understand"}
Difficulty: {difficulty or "Medium"}
Output language: {output_language or "en"}

JSON shape:
{{"questions":[{{"type":"{question_type}","question_text":"Which rule prevents routing loops?","options":{example_options},"correct_answer":"{example_answer}","explanation":"Supported by the slide.","bloom_level":"{bloom_level or "Understand"}","difficulty":"{difficulty or "Medium"}","source_slide":null}}]}}
""".strip()

    return f"""
You are generating instructor-reviewed classroom engagement questions.
Return only valid JSON. Do not include markdown fences.
Generate classroom engagement questions, not exam-style copy questions.
Use the lecture content as source, but rewrite in instructor-friendly wording.
Keep wording concise and use simple academic language.
Use only concrete facts from the provided slide excerpts.
Never ask "What is the lecture content?", "What is this lecture about?", or any broad summary question.
Ask about a specific decision, consequence, comparison, problem, solution, tradeoff, or scenario from a slide.
Avoid simple definition questions.
Avoid questions based only on slide titles; use the slide details and relationships between ideas.
Prefer understanding and application questions over recall-only questions.
Question text must be 25 words or fewer.
For multiple choice questions, use exactly 4 answer options.
For multiple choice questions, each option must be 15 words or fewer.
For multiple choice questions, options must be short balanced statements, not paragraphs.
For multiple choice questions, do not paste full slide sentences into options.
For multiple choice questions, make distractors plausible, related to the same concept, and not random excerpts from other slides.
For multiple choice questions, avoid obvious wrong answers.
For multiple choice questions, question_text should be direct, such as "Which statement correctly describes a relation schema?"
Do not use generic wording like "Which option best explains the impact of this slide's idea?"
For the short answer question, ask one clear idea in 25 words or fewer.
For the short answer question, correct_answer must be no more than 14 words.
Do not make the short answer require a paragraph, long explanation, or essay-style response.
For the short answer question, never include the correct answer, answer phrase, or a near-copy of it in question_text.
For the short answer question, do not write "Example:" in question_text or copy example text from the slide into question_text.
Include source_slide when a slide number is available.

Lecture content:
{base}
{regenerate}

Create exactly:
- 3 multiple choice questions
- 1 short answer question

JSON shape:
{{
  "questions": [
    {{
      "type": "mcq",
      "question_text": "",
      "options": ["", "", "", ""],
      "correct_answer": "",
      "explanation": "",
      "bloom_level": "",
      "difficulty": "",
      "source_slide": null
    }},
    {{
      "type": "short_answer",
      "question_text": "",
      "options": [],
      "correct_answer": "",
      "explanation": "",
      "bloom_level": "",
      "difficulty": "",
      "source_slide": null
    }}
  ]
}}
""".strip()


def build_compact_retry_question_prompt(
    text: str,
    question_type: str,
    bloom_level: str | None = None,
    difficulty: str | None = None,
    output_language: str | None = None,
    question_index: int | None = None,
    avoid_questions: list[str] | None = None,
) -> str:
    settings = get_settings()
    base = build_prompt_context(text, min(settings.ollama_prompt_chars, 700))
    avoid = "\n".join(f"- {question}" for question in (avoid_questions or []) if question)
    avoid_block = f"\nAvoid these previous questions:\n{avoid}\n" if avoid else ""
    if question_type == "mcq":
        schema = (
            '{"questions":[{"type":"mcq","question_text":"...","options":["...","...","...","..."],'
            '"correct_answer":"...","explanation":"...","bloom_level":"'
            f'{bloom_level or "Understand"}","difficulty":"{difficulty or "Medium"}","source_slide":null}}]}}'
        )
        type_rules = """
MCQ hard limits:
- question_text: 12 words maximum
- each option: 4 words maximum
- correct_answer must exactly equal one option
- explanation: 6 words maximum
- total response under 90 words
""".strip()
    else:
        schema = (
            '{"questions":[{"type":"short_answer","question_text":"...","options":[],'
            '"correct_answer":"...","explanation":"...","bloom_level":"'
            f'{bloom_level or "Understand"}","difficulty":"{difficulty or "Medium"}","source_slide":null}}]}}'
        )
        type_rules = """
Short-answer hard limits:
- question_text: 12 words maximum
- options must be []
- correct_answer: 8 words maximum
- explanation: 6 words maximum
- total response under 70 words
""".strip()

    return f"""
Return one line of minified JSON only.
Stop immediately after the final closing brace.
Do not include markdown, commentary, labels, or a second object.
Generate exactly one {question_type} question from the lecture excerpt.
Use this exact root and keys: {schema}
{type_rules}
Bloom level: {bloom_level or "Understand"}
Difficulty: {difficulty or "Medium"}
Output language: {output_language or "en"}
Question number: {question_index or 1}
{avoid_block}
Lecture excerpt:
{base}
""".strip()


def extract_first_balanced_json_object(text: str, start_index: int = 0) -> dict | None:
    start = text.find("{", start_index)
    if start == -1:
        return None

    depth = 0
    in_string = False
    escaped = False
    for index in range(start, len(text)):
        char = text[index]
        if escaped:
            escaped = False
            continue
        if char == "\\":
            escaped = True
            continue
        if char == '"':
            in_string = not in_string
            continue
        if in_string:
            continue
        if char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                try:
                    parsed = json.loads(text[start : index + 1])
                except json.JSONDecodeError:
                    return None
                return parsed if isinstance(parsed, dict) else None
    return None


def parse_ollama_question_payload(model_response: str | dict | list) -> dict | list:
    if not isinstance(model_response, str):
        return model_response

    cleaned = model_response.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", cleaned, flags=re.IGNORECASE | re.DOTALL).strip()

    try:
        return json.loads(cleaned)
    except json.JSONDecodeError as exc:
        questions_match = re.search(r'"questions"\s*:\s*\[', cleaned)
        if questions_match:
            first_question = extract_first_balanced_json_object(cleaned, questions_match.end())
            if first_question:
                return {"questions": [first_question]}

        first_object = extract_first_balanced_json_object(cleaned)
        if first_object and "questions" not in first_object:
            return {"questions": [first_object]}

        preview = cleaned[:300].replace("\n", " ")
        raise RuntimeError(f"Ollama returned invalid JSON for questions. Preview: {preview}") from exc


def fallback_question_from_lecture(
    text: str,
    question_type: str | None,
    bloom_level: str | None,
    difficulty: str | None,
    question_index: int | None = None,
    avoid_questions: list[str] | None = None,
) -> dict:
    return local_clean_question_from_lecture(text, question_type, bloom_level, difficulty, question_index, avoid_questions)


def fallback_questions_from_lecture(
    text: str,
    question_type: str | None,
    bloom_level: str | None,
    difficulty: str | None,
    question_index: int | None = None,
    avoid_questions: list[str] | None = None,
) -> list[InstructorQuestion]:
    question = normalize_question(
        fallback_question_from_lecture(
            text,
            question_type,
            bloom_level,
            difficulty,
            question_index,
            avoid_questions,
        )
    )
    if question_type:
        question.type = question_type
        question.bloom_level = bloom_level or question.bloom_level
        question.difficulty = difficulty or question.difficulty
        if question_type == "short_answer":
            question.options = []
    return [question]


def repair_question_dict(
    raw: dict,
    text: str,
    question_type: str | None,
    bloom_level: str | None,
    difficulty: str | None,
    question_index: int | None = None,
    avoid_questions: list[str] | None = None,
) -> dict:
    fallback = fallback_question_from_lecture(text, question_type or raw.get("type"), bloom_level, difficulty, question_index, avoid_questions)
    repaired = {**fallback, **raw}

    placeholder_question = str(repaired.get("question_text") or "").strip().lower()
    if (
        not repaired.get("question_text")
        or placeholder_question.startswith("a complete question based on the lecture content")
        or is_weak_question_text(placeholder_question, avoid_questions)
    ):
        repaired["question_text"] = fallback["question_text"]
    if not repaired.get("correct_answer") or repaired.get("correct_answer") == "Definition":
        repaired["correct_answer"] = fallback["correct_answer"]
    if not repaired.get("explanation") and raw.get("explanaition"):
        repaired["explanation"] = raw.get("explanaition")
    if not repaired.get("explanation"):
        repaired["explanation"] = fallback["explanation"]

    if (question_type or repaired.get("type")) == "mcq":
        options = [option for option in repaired.get("options", []) if str(option).strip()]
        if options == ["Definition", "Example", "Process", "Outcome"]:
            options = []
        fallback_options = fallback["options"]
        repaired["options"] = (options + fallback_options)[:4]
        if repaired["correct_answer"] not in repaired["options"]:
            repaired["correct_answer"] = fallback["correct_answer"]
        repaired["options"] = shuffled_mcq_options(repaired["options"], repaired["correct_answer"])
    else:
        repaired["options"] = []
        repaired["correct_answer"] = concise_short_answer(repaired.get("correct_answer") or fallback["correct_answer"])
        if "example:" in str(repaired.get("question_text", "")).lower() or answer_leaks_into_question(
            repaired.get("question_text", ""),
            repaired.get("correct_answer", ""),
        ):
            repaired["question_text"] = fallback["question_text"]
        if not re.search(r"few words|short sentence", str(repaired.get("question_text", "")), flags=re.IGNORECASE):
            repaired["question_text"] = f"In a few words or one short sentence, {str(repaired['question_text']).strip()}"
        if answer_leaks_into_question(repaired.get("question_text", ""), repaired.get("correct_answer", "")):
            repaired["question_text"] = "In a few words or one short sentence, what is the key conclusion from this slide?"

    return repaired


def strict_quality_suffix() -> str:
    return """

STRICT QUALITY PASS:
Your previous output failed classroom-quality validation.
Rewrite the question from the same concept using these hard limits:
- question_text: 25 words or fewer
- MCQ options: exactly 4, each 15 words or fewer
- no option may copy a full slide sentence
- no multi-sentence options
- no generic wording about "this slide's idea"
- distractors must be plausible misconceptions about the same concept
- short-answer questions must ask one direct why/what/how question
Return only the same JSON shape.
""".rstrip()


def question_dicts_from_ollama_raw(
    raw: dict,
    text: str,
    question_type: str | None,
    bloom_level: str | None,
    difficulty: str | None,
    question_index: int | None = None,
    avoid_questions: list[str] | None = None,
) -> list[dict]:
    model_response = raw.get("response", raw)
    parsed = parse_ollama_question_payload(model_response)

    questions = parsed.get("questions", parsed if isinstance(parsed, list) else [])
    if not isinstance(questions, list):
        return []
    return [question for question in questions if isinstance(question, dict)]


def repair_and_validate_question_dicts(
    questions: list[dict],
    text: str,
    question_type: str | None,
    bloom_level: str | None,
    difficulty: str | None,
    question_index: int | None = None,
    avoid_questions: list[str] | None = None,
) -> tuple[list[dict], bool]:
    cleaned_questions: list[dict] = []
    had_quality_failure = False

    for question in questions:
        target_type = question_type or question.get("type")
        fallback = fallback_question_from_lecture(text, target_type, bloom_level, difficulty, question_index, avoid_questions)
        repaired = repair_question_dict(question, text, target_type, bloom_level, difficulty, question_index, avoid_questions)
        repaired, errors = enforce_question_quality(repaired, text, fallback)
        if errors:
            had_quality_failure = True
            logger.info("Question failed quality validation (%s); using clean local fallback.", ", ".join(errors))
            repaired = fallback
        cleaned_questions.append(repaired)

    return cleaned_questions, had_quality_failure


def call_ollama_for_questions(
    text: str,
    regenerate_question: InstructorQuestion | None = None,
    question_type: str | None = None,
    bloom_level: str | None = None,
    difficulty: str | None = None,
    output_language: str | None = None,
    question_index: int | None = None,
    avoid_questions: list[str] | None = None,
    debug_lecture_id: str | None = None,
    debug_class_id: str | None = None,
    debug_instructor_id: str | None = None,
    debug_batch_id: str | None = None,
    debug_question_number: int | None = None,
    debug_question_type: str | None = None,
    debug_request_kind: str = "initial",
    debug_parent_request_id: str | None = None,
    llm_provider: str | None = None,
    llm_model: str | None = None,
) -> list[InstructorQuestion]:
    settings = get_settings()
    provider_config = resolve_ollama_provider(llm_provider, llm_model)
    options = {
        "temperature": 1,
        "num_ctx": settings.ollama_num_ctx,
        "num_predict": settings.ollama_num_predict,
        "num_gpu": settings.ollama_num_gpu,
    }
    if settings.ollama_num_thread > 0:
        options["num_thread"] = settings.ollama_num_thread

    def export_request_debug(
        *,
        timestamp: datetime,
        prompt: str,
        payload: dict,
        raw_response: str | None,
        parsed_response: dict | list | str | None,
        generation_success: bool,
        processing_time_ms: int | None,
        http_status: int | None,
        request_metadata: dict,
        error: dict | None,
    ) -> None:
        write_llm_debug_export(
            timestamp=timestamp,
            lecture_id=debug_lecture_id,
            class_id=debug_class_id,
            instructor_id=debug_instructor_id,
            model=provider_config.model,
            provider=provider_config.provider,
            system_prompt="",
            user_prompt=prompt,
            request_payload=payload,
            raw_response=raw_response,
            parsed_ollama_response=parsed_response,
            generation_success=generation_success,
            processing_time_ms=processing_time_ms,
            http_status=http_status,
            request_metadata=request_metadata,
            input_metadata={
                "question_type": debug_question_type or question_type,
                "question_number": debug_question_number or question_index,
                "bloom_level": bloom_level,
                "difficulty": difficulty,
                "output_language": output_language,
                "requested_question_count": 1 if (debug_question_type or question_type) else 4,
                "previous_questions": avoid_questions or [],
                "provider": provider_config.provider,
                "deployment_type": provider_config.deployment_type,
                "requested_model": provider_config.model,
                "endpoint_label": provider_config.endpoint_label,
            },
            lecture_text=text,
            error=error,
        )

    def request_questions(
        prompt: str,
        attempts: int = 2,
        request_kind: str | None = None,
        parent_request_id: str | None = None,
        starting_attempt_number: int = 1,
    ) -> tuple[dict | None, RuntimeError | None, dict | None]:
        last_error: RuntimeError | None = None
        previous_request_id = parent_request_id
        last_request_metadata: dict | None = None
        for attempt in range(attempts):
            request_timestamp = datetime.now(timezone.utc)
            current_kind = request_kind or debug_request_kind or "initial"
            if attempt > 0:
                current_kind = "automatic_retry"
            request_metadata = {
                "request_id": new_debug_uuid(),
                "batch_id": debug_batch_id,
                "question_number": debug_question_number,
                "question_type": debug_question_type,
                "attempt_number": starting_attempt_number + attempt,
                "request_kind": current_kind,
                "is_retry": attempt > 0 or current_kind == "quality_retry",
                "is_regeneration": current_kind == "manual_regeneration",
                "parent_request_id": previous_request_id if attempt > 0 or current_kind in {"quality_retry", "manual_regeneration"} else None,
                "provider": provider_config.provider,
                "deployment_type": provider_config.deployment_type,
                "requested_model": provider_config.model,
                "endpoint_label": provider_config.endpoint_label,
            }
            last_request_metadata = request_metadata
            result = perform_ollama_generate_request(
                provider=provider_config.provider,
                base_url=provider_config.base_url,
                model=provider_config.model,
                api_key=provider_config.api_key,
                prompt=prompt,
                options=options,
                timeout=settings.ollama_timeout_seconds,
                endpoint_label=provider_config.endpoint_label,
            )
            payload = result["safe_payload"]
            parsed_raw = result.get("parsed_response")
            raw_response_text = result.get("raw_response_text")
            model_response_text = parsed_raw.get("response", raw_response_text) if isinstance(parsed_raw, dict) else raw_response_text
            response_metadata = provider_config_metadata(provider_config, parsed_raw, result.get("http_status"), result.get("elapsed_ms"))
            request_metadata.update(response_metadata)
            if result.get("error") is None and isinstance(parsed_raw, dict) and parsed_raw.get("done_reason") == "length":
                result["error"] = "Ollama output reached num_predict before completing JSON."
                result["error_category"] = "truncated_output"
            elif result.get("error") is None and isinstance(parsed_raw, dict) and parsed_raw.get("done") is True:
                export_request_debug(
                    timestamp=request_timestamp,
                    prompt=prompt,
                    payload=payload,
                    raw_response=model_response_text,
                    parsed_response=parsed_raw,
                    generation_success=True,
                    processing_time_ms=result.get("elapsed_ms"),
                    http_status=result.get("http_status"),
                    request_metadata=request_metadata,
                    error=None,
                )
                return parsed_raw, None, request_metadata

            if result.get("error") is None and isinstance(parsed_raw, dict) and parsed_raw.get("done") is False:
                result["error"] = "Ollama returned done=false for a non-streaming generation response."
                result["error_category"] = "incomplete_response"

            error_category = result.get("error_category") or "request_failed"
            if error_category == "timeout":
                error_message = (
                    f"Ollama timed out after {settings.ollama_timeout_seconds}s while generating questions. "
                    "The model may still be loading or running slowly on CPU. Try again once it is warm, "
                    f"or use a smaller model / increase OLLAMA_TIMEOUT_SECONDS."
                )
            elif error_category == "local_ollama_not_running":
                error_message = f"Local Ollama is not reachable. Start Ollama and pull {provider_config.model}."
            elif error_category == "endpoint_unreachable":
                error_message = "Cloud Ollama endpoint is unreachable."
            elif error_category == "authentication_failed":
                error_message = "Ollama cloud authentication failed."
            elif error_category == "model_not_found":
                error_message = f"Configured Ollama model is unavailable: {provider_config.model}."
            elif error_category == "rate_limited":
                error_message = "Ollama request was rate limited (HTTP 429)."
            elif error_category == "server_error":
                error_message = f"Ollama returned a server error (HTTP {result.get('http_status')})."
            elif error_category == "malformed_json":
                error_message = "Ollama returned a malformed JSON response."
            elif error_category == "incomplete_response":
                error_message = "Ollama returned done=false for a non-streaming generation response."
            elif error_category == "truncated_output":
                error_message = (
                    "Ollama output was truncated before valid question JSON was complete "
                    f"(done_reason=length, completion_tokens={request_metadata.get('completion_tokens')}). "
                    "Increase OLLAMA_NUM_PREDICT or ask for a shorter question."
                )
            else:
                error_message = result.get("error") or "Ollama request failed."

            request_metadata["error_category"] = error_category
            request_metadata["error_message"] = redact_secret_text(error_message)
            if error_category in {"truncated_output", "malformed_json", "incomplete_response"} and isinstance(parsed_raw, dict):
                logger.warning(
                    "Ollama generation output failure provider=%s model=%s done=%s done_reason=%s prompt_tokens=%s completion_tokens=%s http_status=%s preview=%s",
                    provider_config.provider,
                    provider_config.model,
                    parsed_raw.get("done"),
                    parsed_raw.get("done_reason"),
                    request_metadata.get("prompt_tokens"),
                    request_metadata.get("completion_tokens"),
                    result.get("http_status"),
                    model_response_preview(parsed_raw),
                )

            last_error = RuntimeError(redact_secret_text(error_message))
            exc = result.get("exception")
            if isinstance(exc, BaseException):
                last_error.__cause__ = exc
            export_request_debug(
                timestamp=request_timestamp,
                prompt=prompt,
                payload=payload,
                raw_response=model_response_text or raw_response_text,
                parsed_response=parsed_raw,
                generation_success=False,
                processing_time_ms=result.get("elapsed_ms"),
                http_status=result.get("http_status"),
                request_metadata=request_metadata,
                error={
                    "type": error_category,
                    "message": str(last_error),
                    "traceback": "" if exc is None else redact_secret_text("".join(traceback.format_exception(exc))),
                },
            )

            if attempt == 0 and attempts > 1 and error_category != "truncated_output":
                previous_request_id = request_metadata["request_id"]
                time.sleep(1)
                continue
            break
        return None, last_error, last_request_metadata

    prompt = build_question_prompt(
        text,
        regenerate_question,
        question_type,
        bloom_level,
        difficulty,
        output_language,
        question_index,
        avoid_questions,
    )
    raw, last_error, source_request_metadata = request_questions(prompt, parent_request_id=debug_parent_request_id)

    if raw is None:
        if question_type and (source_request_metadata or {}).get("error_category") == "truncated_output":
            compact_prompt = build_compact_retry_question_prompt(
                text,
                question_type,
                bloom_level,
                difficulty,
                output_language,
                question_index,
                avoid_questions,
            )
            logger.info(
                "Retrying Ollama generation with compact prompt after truncation provider=%s model=%s question_type=%s",
                provider_config.provider,
                provider_config.model,
                question_type,
            )
            compact_raw, compact_error, compact_request_metadata = request_questions(
                compact_prompt,
                attempts=1,
                request_kind="compact_truncation_retry",
                parent_request_id=(source_request_metadata or {}).get("request_id"),
                starting_attempt_number=((source_request_metadata or {}).get("attempt_number") or 1) + 1,
            )
            if compact_raw is not None:
                raw = compact_raw
                source_request_metadata = compact_request_metadata
            else:
                raise compact_error or last_error or RuntimeError("Ollama request failed after compact retry.")
        else:
            raise last_error or RuntimeError("Ollama request failed.")

    parse_error: RuntimeError | None = None
    try:
        raw_questions = question_dicts_from_ollama_raw(raw, text, question_type, bloom_level, difficulty, question_index, avoid_questions)
    except RuntimeError as exc:
        parse_error = exc
        raw_questions = []

    if parse_error and question_type and (source_request_metadata or {}).get("request_kind") != "compact_truncation_retry":
        compact_prompt = build_compact_retry_question_prompt(
            text,
            question_type,
            bloom_level,
            difficulty,
            output_language,
            question_index,
            avoid_questions,
        )
        logger.info(
            "Retrying Ollama generation with compact prompt after malformed JSON provider=%s model=%s question_type=%s",
            provider_config.provider,
            provider_config.model,
            question_type,
        )
        compact_raw, compact_error, compact_request_metadata = request_questions(
            compact_prompt,
            attempts=1,
            request_kind="compact_truncation_retry",
            parent_request_id=(source_request_metadata or {}).get("request_id"),
            starting_attempt_number=((source_request_metadata or {}).get("attempt_number") or 1) + 1,
        )
        if compact_raw is not None:
            raw = compact_raw
            source_request_metadata = compact_request_metadata
            try:
                raw_questions = question_dicts_from_ollama_raw(raw, text, question_type, bloom_level, difficulty, question_index, avoid_questions)
                parse_error = None
            except RuntimeError as exc:
                parse_error = exc
        else:
            parse_error = compact_error or parse_error

    if parse_error:
        preview = model_response_preview(raw)
        logger.warning(
            "Ollama returned malformed question JSON provider=%s model=%s done=%s done_reason=%s prompt_tokens=%s completion_tokens=%s http_status=%s preview=%s",
            provider_config.provider,
            provider_config.model,
            raw.get("done") if isinstance(raw, dict) else None,
            raw.get("done_reason") if isinstance(raw, dict) else None,
            raw.get("prompt_eval_count") if isinstance(raw, dict) else None,
            raw.get("eval_count") if isinstance(raw, dict) else None,
            (source_request_metadata or {}).get("http_status"),
            preview,
        )
        raise RuntimeError(
            "Ollama returned malformed question JSON "
            f"(done_reason={(raw or {}).get('done_reason') if isinstance(raw, dict) else None}, "
            f"completion_tokens={(raw or {}).get('eval_count') if isinstance(raw, dict) else None}). "
            f"Preview: {preview}"
        ) from parse_error
    questions, had_quality_failure = repair_and_validate_question_dicts(
        raw_questions,
        text,
        question_type,
        bloom_level,
        difficulty,
        question_index,
        avoid_questions,
    )
    if had_quality_failure:
        strict_raw, strict_error, strict_request_metadata = request_questions(
            prompt + strict_quality_suffix(),
            attempts=1,
            request_kind="quality_retry",
            parent_request_id=(source_request_metadata or {}).get("request_id"),
            starting_attempt_number=((source_request_metadata or {}).get("attempt_number") or 1) + 1,
        )
        if strict_raw is not None:
            strict_questions = question_dicts_from_ollama_raw(strict_raw, text, question_type, bloom_level, difficulty, question_index, avoid_questions)
            strict_cleaned, _strict_failed = repair_and_validate_question_dicts(
                strict_questions,
                text,
                question_type,
                bloom_level,
                difficulty,
                question_index,
                avoid_questions,
            )
            if strict_cleaned:
                questions = strict_cleaned
                source_request_metadata = strict_request_metadata
        elif strict_error:
            logger.info("Strict question regeneration failed; using clean local fallback: %s", strict_error)

    if not questions:
        questions = [fallback_question_from_lecture(text, question_type, bloom_level, difficulty, question_index, avoid_questions)]
    normalized = [normalize_question(question) for question in questions]
    if question_type:
        normalized = normalized[:1]
        for question in normalized:
            question.type = question_type
            question.bloom_level = bloom_level or question.bloom_level
            question.difficulty = difficulty or question.difficulty
            if question_type == "short_answer":
                question.options = []
    for question in normalized:
        question.generation_metadata = generated_question_metadata(source_request_metadata, provider_config)
        remember_question_debug_request(question.question_id, source_request_metadata)
    return normalized


def normalize_question(raw: dict) -> InstructorQuestion:
    question_type = raw.get("type", "mcq")
    if question_type not in {"mcq", "short_answer"}:
        question_type = "short_answer" if "short" in question_type else "mcq"

    options = raw.get("options") or []
    correct_answer = raw.get("correct_answer") or ""
    if question_type == "short_answer":
        options = []
        correct_answer = concise_short_answer(correct_answer)
    else:
        options = shuffled_mcq_options(options, correct_answer)

    return InstructorQuestion(
        question_id=raw.get("question_id") or new_id("question"),
        upload_id=raw.get("upload_id"),
        type=question_type,
        question_text=raw.get("question_text") or raw.get("prompt") or "Untitled question",
        options=options,
        correct_answer=correct_answer,
        explanation=raw.get("explanation") or raw.get("explanaition") or "",
        bloom_level=raw.get("bloom_level") or raw.get("bloom_taxonomy") or "Understanding",
        difficulty=raw.get("difficulty") or "Medium",
        source_slide=raw.get("source_slide"),
        status=raw.get("status") or "generated",
    )


def make_session_code() -> str:
    return "".join(random.choices(string.ascii_uppercase + string.digits, k=6))


def make_join_link(session_code: str) -> str:
    frontend_origin = get_settings().frontend_origin.rstrip("/")
    return f"{frontend_origin}/student/join/{session_code}"


def make_qr_base64(data: str) -> str:
    image = qrcode.make(data)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
    return f"data:image/png;base64,{encoded}"


def get_upload_file_path(upload_id: str, file_type: str | None = None) -> Path | None:
    upload_dir = get_storage_root() / "uploads"
    suffixes = [f".{file_type.lower().lstrip('.')}"] if file_type else [".pptx", ".pdf"]
    for suffix in suffixes:
        candidate = upload_dir / f"{upload_id}{suffix}"
        if candidate.exists():
            return candidate
    return None


def add_slide_title(slide, text: str) -> None:
    if slide.shapes.title:
        slide.shapes.title.text = text
        return
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(8.8), Inches(0.6))
    paragraph = title_box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True


def style_shape(shape, fill: RGBColor, line: RGBColor | None = None) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill


def set_text_box(
    shape,
    text: str,
    *,
    size: int,
    color: RGBColor,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = Inches(0.18)
    frame.margin_right = Inches(0.18)
    frame.margin_top = Inches(0.08)
    frame.margin_bottom = Inches(0.08)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color


def move_last_slide_to_index(presentation: Presentation, target_index: int) -> None:
    sld_id_list = presentation.slides._sldIdLst
    last_slide_id = sld_id_list[-1]
    sld_id_list.remove(last_slide_id)
    safe_index = max(0, min(target_index, len(sld_id_list)))
    sld_id_list.insert(safe_index, last_slide_id)


def add_join_slide(presentation: Presentation, session_code: str | None = None, join_link: str | None = None) -> None:
    # Mirror the look-and-feel of question slides so the deck stays consistent.
    slide_width = presentation.slide_width
    slide_height = presentation.slide_height
    margin_x = Inches(0.8)
    content_width = slide_width - (margin_x * 2)
    teal = RGBColor(22, 119, 132)
    deep_teal = RGBColor(8, 84, 97)
    mint = RGBColor(222, 247, 235)
    soft_bg = RGBColor(246, 249, 250)
    white = RGBColor(255, 255, 255)
    slate = RGBColor(71, 85, 105)
    border = RGBColor(212, 226, 229)

    layout = presentation.slide_layouts[6] if len(presentation.slide_layouts) > 6 else presentation.slide_layouts[-1]
    slide = presentation.slides.add_slide(layout)
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, slide_width, slide_height)
    style_shape(background, soft_bg)

    header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, margin_x, Inches(0.45), Inches(3.0), Inches(0.42))
    style_shape(header, mint)
    set_text_box(header, "JOIN SESSION", size=12, color=deep_teal, bold=True, align=PP_ALIGN.CENTER)

    gap = Inches(0.3)
    qr_width = Inches(2.2)
    code_width = content_width - qr_width - gap

    code_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, margin_x, Inches(1.15), code_width, Inches(2.4))
    style_shape(code_box, white, border)
    if session_code and not join_link:
        join_link = make_join_link(session_code)
    code_text = (
        f"Session code: {session_code or '________'}\n\nJoin link:\n{join_link or '________'}\n"
    )
    set_text_box(code_box, code_text, size=18, color=RGBColor(15, 23, 42), bold=True, align=PP_ALIGN.LEFT)

    qr_x = margin_x + code_width + gap
    qr_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, qr_x, Inches(1.15), qr_width, qr_width)
    style_shape(qr_box, white, border)
    if join_link:
        qr_image = qrcode.make(join_link)
        buffer = io.BytesIO()
        qr_image.save(buffer, format="PNG")
        buffer.seek(0)
        slide.shapes.add_picture(buffer, qr_x + Inches(0.12), Inches(1.27), width=qr_width - Inches(0.24), height=qr_width - Inches(0.24))
    else:
        set_text_box(qr_box, "QR code\nplaceholder", size=16, color=teal, bold=True, align=PP_ALIGN.CENTER)

    #footer = slide.shapes.add_textbox(margin_x, slide_height - Inches(0.65), content_width, Inches(0.28))
    #set_text_box(footer, "Add the join code and QR code manually or use the instructor dashboard to download the QR.", size=11, color=slate, align=PP_ALIGN.CENTER)

    # Move this join slide to be right after the title (index 1).
    try:
        move_last_slide_to_index(presentation, 1)
    except Exception:
        # best-effort; if manipulation fails leave slide as appended
        pass


def add_question_slide(presentation: Presentation, question: InstructorQuestion, index: int) -> None:
    question_layout = presentation.slide_layouts[6] if len(presentation.slide_layouts) > 6 else presentation.slide_layouts[-1]
    slide_width = presentation.slide_width
    slide_height = presentation.slide_height
    margin_x = Inches(0.8)
    content_width = slide_width - (margin_x * 2)
    teal = RGBColor(22, 119, 132)
    deep_teal = RGBColor(8, 84, 97)
    mint = RGBColor(222, 247, 235)
    soft_bg = RGBColor(246, 249, 250)
    white = RGBColor(255, 255, 255)
    slate = RGBColor(71, 85, 105)
    border = RGBColor(212, 226, 229)

    slide = presentation.slides.add_slide(question_layout)
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, slide_width, slide_height)
    style_shape(background, soft_bg)

    header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, margin_x, Inches(0.45), Inches(3.0), Inches(0.42))
    style_shape(header, mint)
    set_text_box(header, f"QUESTION {index}", size=12, color=deep_teal, bold=True, align=PP_ALIGN.CENTER)

    meta = slide.shapes.add_textbox(slide_width - margin_x - Inches(3.2), Inches(0.45), Inches(3.2), Inches(0.42))
    set_text_box(meta, f"{question.bloom_level} | {question.difficulty}", size=11, color=slate, bold=True, align=PP_ALIGN.RIGHT)

    question_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, margin_x, Inches(1.15), content_width, Inches(1.45))
    style_shape(question_card, white, border)
    set_text_box(question_card, question.question_text, size=24, color=RGBColor(15, 23, 42), bold=True)

    options = question.options[:6]
    if options:
        columns = 2 if len(options) > 2 else 1
        gap = Inches(0.22)
        option_width = (content_width - (gap * (columns - 1))) // columns
        option_height = Inches(0.78)
        start_y = Inches(2.95)

        for option_index, option in enumerate(options):
            row = option_index // columns
            column = option_index % columns
            x = margin_x + (column * (option_width + gap))
            y = start_y + (row * (option_height + gap))
            option_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, option_width, option_height)
            style_shape(option_card, white, border)
            letter = chr(65 + option_index)
            set_text_box(option_card, f"{letter}. {option}", size=16, color=RGBColor(30, 41, 59), bold=True)
    else:
        response_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, margin_x, Inches(3.05), content_width, Inches(1.05))
        style_shape(response_card, white, border)
        set_text_box(response_card, "Write your short answer", size=18, color=teal, bold=True, align=PP_ALIGN.CENTER)

    footer = slide.shapes.add_textbox(margin_x, slide_height - Inches(0.65), content_width, Inches(0.28))
    set_text_box(footer, "Discuss, answer, then reveal feedback in the live session.", size=11, color=slate, align=PP_ALIGN.CENTER)


def add_question_slides(presentation: Presentation, questions: list[InstructorQuestion]) -> None:
    for index, question in enumerate(questions, start=1):
        add_question_slide(presentation, question, index)


QUESTION_MATCH_STOP_WORDS = GENERIC_LECTURE_TERMS | {
    "answer",
    "based",
    "best",
    "correct",
    "difficulty",
    "engagement",
    "explain",
    "idea",
    "important",
    "lecture",
    "level",
    "main",
    "option",
    "original",
    "question",
    "recent",
    "review",
    "section",
    "support",
    "topic",
    "understand",
}


def extract_slide_text(slide) -> str:
    lines: list[str] = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            lines.append(re.sub(r"\s+", " ", shape.text).strip())
    return "\n".join(lines)


def presentation_slide_texts(presentation: Presentation, slide_count: int) -> dict[int, str]:
    slides = list(presentation.slides)
    if len(slides) > slide_count and slide_count > 1:
        original_slides = [slides[0], *slides[2 : slide_count + 1]]
    else:
        original_slides = slides[:slide_count]
    return {
        slide_index: extract_slide_text(slide)
        for slide_index, slide in enumerate(original_slides, start=1)
    }


def is_content_presentation_slide(text: str) -> bool:
    lines = [re.sub(r"\s+", " ", line).strip(" -:\u2022") for line in text.splitlines()]
    lines = [line for line in lines if line and not is_low_value_slide_line(line)]
    if not lines:
        return False
    combined = " ".join(lines)
    if len(combined) < 35:
        return False
    section = {"title": lines[0], "body": lines[1:], "lines": lines}
    return len(lines) == 1 or not is_low_value_slide_section(section)


def content_slide_numbers(slide_texts: dict[int, str]) -> list[int]:
    content_slides = [
        slide_number
        for slide_number, text in slide_texts.items()
        if is_content_presentation_slide(text)
    ]
    if content_slides:
        return content_slides
    fallback = [slide_number for slide_number in slide_texts if slide_number > 1]
    return fallback or list(slide_texts)


def terms_for_matching(text: str) -> set[str]:
    terms = set()
    for term in re.findall(r"[A-Za-z][A-Za-z0-9-]{3,}", text.lower()):
        if term not in QUESTION_MATCH_STOP_WORDS:
            terms.add(term)
    return terms


def question_match_text(question: InstructorQuestion) -> str:
    return " ".join(
        [
            question.question_text or "",
            question.correct_answer or "",
            question.explanation or "",
            " ".join(question.options or []),
        ]
    )


def best_matching_content_slide(
    question: InstructorQuestion,
    slide_texts: dict[int, str],
    content_slides: list[int],
) -> int | None:
    question_terms = terms_for_matching(question_match_text(question))
    if not question_terms:
        return None

    best_slide = None
    best_score = 0
    for slide_number in content_slides:
        slide_terms = terms_for_matching(slide_texts.get(slide_number, ""))
        score = len(question_terms & slide_terms)
        if score > best_score:
            best_slide = slide_number
            best_score = score

    return best_slide if best_score >= 1 else None


def coerce_source_slide(question: InstructorQuestion, original_slide_count: int) -> int | None:
    try:
        source_slide = int(question.source_slide) if question.source_slide is not None else None
    except (TypeError, ValueError):
        return None
    if source_slide is None:
        return None
    return max(1, min(source_slide, original_slide_count))


def cadence_anchor_slide(missing_index: int, missing_count: int, content_slides: list[int]) -> int:
    if not content_slides:
        return 1
    interval = round(len(content_slides) / max(missing_count, 1))
    interval = max(2, min(4, interval or 3))
    content_index = min(len(content_slides) - 1, ((missing_index + 1) * interval) - 1)
    return content_slides[content_index]


def content_position_for_slide(slide_number: int, content_slides: list[int]) -> int:
    if not content_slides:
        return 0
    if slide_number in content_slides:
        return content_slides.index(slide_number)
    for index, content_slide in enumerate(content_slides):
        if content_slide > slide_number:
            return index
    return len(content_slides) - 1


def choose_spread_anchor(preferred_slide: int, previous_position: int | None, used_positions: set[int], content_slides: list[int]) -> int:
    if not content_slides:
        return preferred_slide

    preferred_position = content_position_for_slide(preferred_slide, content_slides)
    min_position = 0 if previous_position is None else previous_position + 2
    target_position = max(preferred_position, min_position)

    # Prefer a nearby future content slide so engagement appears every 2-4 slides instead of clustering.
    future_candidates = [
        position
        for position in range(target_position, min(len(content_slides), target_position + 3))
        if position not in used_positions
    ]
    if future_candidates:
        return content_slides[future_candidates[0]]

    # If there is no future room, pick the closest unused content slide after the preferred slide.
    fallback_candidates = [
        position
        for position in range(preferred_position, len(content_slides))
        if position not in used_positions
    ]
    if fallback_candidates:
        return content_slides[fallback_candidates[0]]

    # Very short decks may not have enough content slides; allow clustering only when unavoidable.
    return content_slides[preferred_position]


def spread_question_assignments(
    assignments: list[tuple[int, int, InstructorQuestion]],
    content_slides: list[int],
) -> list[tuple[int, int, InstructorQuestion]]:
    if len(assignments) <= 1 or not content_slides:
        return assignments

    spread: list[tuple[int, int, InstructorQuestion]] = []
    used_positions: set[int] = set()
    previous_position: int | None = None

    for preferred_slide, order, question in sorted(assignments, key=lambda item: (item[0], item[1])):
        anchor_slide = choose_spread_anchor(preferred_slide, previous_position, used_positions, content_slides)
        anchor_position = content_position_for_slide(anchor_slide, content_slides)
        used_positions.add(anchor_position)
        previous_position = anchor_position
        spread.append((anchor_slide, order, question))

    return sorted(spread, key=lambda item: (item[0], item[1]))


def question_anchor_assignments(
    presentation: Presentation,
    questions: list[InstructorQuestion],
    original_slide_count: int,
) -> list[tuple[int, int, InstructorQuestion]]:
    slide_texts = presentation_slide_texts(presentation, original_slide_count)
    content_slides = content_slide_numbers(slide_texts)
    assignments: list[tuple[int, int, InstructorQuestion]] = []
    missing_source: list[tuple[int, InstructorQuestion]] = []

    for order, question in enumerate(questions):
        source_slide = coerce_source_slide(question, original_slide_count)
        if source_slide:
            assignments.append((source_slide, order, question))
        else:
            missing_source.append((order, question))

    for missing_index, (order, question) in enumerate(missing_source):
        anchor_slide = best_matching_content_slide(question, slide_texts, content_slides)
        if anchor_slide is None:
            anchor_slide = cadence_anchor_slide(missing_index, len(missing_source), content_slides)
        assignments.append((anchor_slide, order, question))

    return spread_question_assignments(assignments, content_slides)


def question_insert_index_after_original_slide(source_slide: int, original_slide_count: int) -> int:
    source_slide = max(1, min(source_slide, original_slide_count))
    original_slide_index = source_slide - 1
    if original_slide_count > 1 and source_slide >= 2:
        original_slide_index += 1
    return max(original_slide_index + 1, 2)


def add_interleaved_question_slides(
    presentation: Presentation,
    questions: list[InstructorQuestion],
    original_slide_count: int,
) -> None:
    assignments = question_anchor_assignments(presentation, questions, original_slide_count)
    inserted_questions = 0
    for question_number, (source_slide, _order, question) in enumerate(assignments, start=1):
        add_question_slide(presentation, question, question_number)
        base_insert_index = question_insert_index_after_original_slide(source_slide, original_slide_count)
        move_last_slide_to_index(presentation, base_insert_index + inserted_questions)
        inserted_questions += 1


def create_question_deck(questions: list[InstructorQuestion], session_code: str | None = None) -> Presentation:
    presentation = Presentation()
    title_layout = presentation.slide_layouts[0]
    title_slide = presentation.slides.add_slide(title_layout)
    add_slide_title(title_slide, "Engagement Questions")
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = "Generated by AILA"
    add_join_slide(presentation, session_code=session_code)
    add_question_slides(presentation, questions)
    return presentation


def reconstruct_presentation(questions: list[InstructorQuestion], upload_id: str, upload: dict | None = None, session_code: str | None = None) -> tuple[str, Path]:
    upload_path = get_upload_file_path(upload_id, upload.get("file_type") if upload else None)
    original_filename = upload.get("filename") if upload else None
    if upload_path and upload_path.suffix.lower() == ".pptx":
        presentation = Presentation(upload_path)
        original_slide_count = len(presentation.slides)
        filename_stem = Path(original_filename or upload_path.name).stem
        filename = f"{filename_stem}_with_questions.pptx"
        add_join_slide(presentation, session_code=session_code)
        add_interleaved_question_slides(presentation, questions, original_slide_count)
    else:
        presentation = create_question_deck(questions, session_code=session_code)
        filename = f"smartclass_engagement_{upload_id}.pptx"

    file_id = new_id("pptx")
    path = get_storage_root() / "presentations" / f"{file_id}.pptx"
    presentation.save(path)
    return filename, path
